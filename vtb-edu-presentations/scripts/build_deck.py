#!/usr/bin/env python3
"""Сборка презентации в стиле «ВТБ Образование» из JSON-описания.

    python3 build_deck.py deck.json -o out/            # HTML + PNG + PDF + PPTX
    python3 build_deck.py deck.json -o out/ --html-only

На выходе в папке out/:
    deck.html          все слайды подряд (просмотр в браузере)
    slide_01.png …     каждый слайд 1920×1080
    deck.pdf           один слайд = одна страница 16:9
    deck.pptx          слайды-картинки 16:9 + заметки докладчика
    report.json        проверка переполнения: какие слайды не влезли

Формат deck.json и список макетов — references/layouts.md.
"""
import argparse
import base64
import html
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

SKILL = Path(__file__).resolve().parent.parent
ASSETS = SKILL / "assets"
W, H = 1920, 1080

# --- Токены бренда (references/brand.md) ----------------------------------
NAVY = "#0A2896"
BLUE = "#1F4FD8"
BLUE2 = "#2E6BFF"
SKY = "#19A3FF"
LIGHT = "#8DB8FF"
PALE = "#DCE6FF"
VIOLET = "#4B1FD6"
VIOLET_DEEP = "#24108A"
CARD = "#F1F3F8"
LINE = "#DDE2EC"
TEXT = "#14161C"
MUTED = "#636B7A"
RED = "#D93A3A"
GREEN = "#1E8E5A"

GRAD_HERO = ("linear-gradient(135deg,#1B0F8C 0%,#3A17D1 38%,#4E2BF0 62%,#2E6BFF 100%)")
GRAD_CARD = f"linear-gradient(100deg,{BLUE} 0%,{BLUE2} 55%,{LIGHT} 100%)"
GRAD_BAR = f"linear-gradient(90deg,{BLUE} 0%,{SKY} 100%)"
GRAD_HEAD = f"linear-gradient(90deg,{NAVY} 0%,{BLUE} 60%,{BLUE2} 100%)"


def data_uri(path):
    p = Path(path)
    mime = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
            ".svg": "image/svg+xml", ".webp": "image/webp", ".ttf": "font/ttf"}[p.suffix.lower()]
    return f"data:{mime};base64," + base64.b64encode(p.read_bytes()).decode()


# --- Типографика: неразрывные пробелы после коротких слов ----------------
_SHORT = re.compile(r"(?<![\w&;])([A-Za-zА-Яа-яЁё]{1,2}|—|–)\s+(?=\S)")


def typo(s):
    """Текст слайда: разрешён инлайн-HTML (<b>, <br>, &nbsp;). Предлоги,
    союзы и тире «приклеиваются» к следующему слову — как в лонгридах курса."""
    if s is None:
        return ""
    s = str(s)
    parts = re.split(r"(<[^>]+>)", s)
    for i in range(0, len(parts), 2):
        parts[i] = _SHORT.sub(lambda m: m.group(1) + "&nbsp;", parts[i])
        parts[i] = re.sub(r"\s+—", "&nbsp;—", parts[i])
    return "".join(parts)


def esc(s):
    return html.escape(str(s), quote=True)


# --- CSS ------------------------------------------------------------------
def css(font_uri):
    return f"""
@font-face {{ font-family:'Inter'; src:url('{font_uri}') format('truetype'); font-weight:100 900; }}
@page {{ size:{W}px {H}px; margin:0; }}
* {{ box-sizing:border-box; margin:0; padding:0; }}
html,body {{ background:#2A2F3A; }}
body {{ font-family:'VTB Group UI','VTB Group','Inter',Arial,sans-serif; color:{TEXT};
  -webkit-font-smoothing:antialiased; font-feature-settings:'tnum' 0; }}
.slide {{ position:relative; width:{W}px; height:{H}px; overflow:hidden; background:#FFFFFF;
  page-break-after:always; break-after:page; }}
.deck-preview .slide {{ margin:0 auto 40px; box-shadow:0 10px 40px rgba(0,0,0,.35); }}

/* фон: тёмный (градиент) и светлый */
.dark {{ background:{GRAD_HERO}; color:#FFFFFF; }}
.dark::before {{ content:''; position:absolute; inset:0;
  background:radial-gradient(circle at 85% 12%,rgba(140,110,255,.55) 0,rgba(140,110,255,0) 42%),
             radial-gradient(circle at 8% 96%,rgba(25,163,255,.35) 0,rgba(25,163,255,0) 38%); }}
.ring {{ position:absolute; border-radius:50%; border:2px solid rgba(255,255,255,.14);
  background:radial-gradient(circle at 30% 30%,rgba(255,255,255,.10),rgba(255,255,255,0) 60%); }}
.light .ring {{ border-color:{PALE}; background:none; }}

/* шапка и подвал */
.logo {{ position:absolute; top:68px; right:96px; height:46px; z-index:5; }}
.kicker {{ position:absolute; top:76px; left:96px; right:640px; z-index:5; font-size:20px; font-weight:600;
  letter-spacing:3px; text-transform:uppercase; color:{BLUE}; white-space:nowrap; overflow:hidden; text-overflow:ellipsis; }}
.dark .kicker {{ color:rgba(255,255,255,.75); }}
.foot {{ position:absolute; left:96px; right:96px; bottom:52px; z-index:5; display:flex;
  justify-content:space-between; align-items:center; font-size:18px; color:{MUTED}; letter-spacing:.5px; }}
.dark .foot {{ color:rgba(255,255,255,.7); }}
.foot .pg {{ font-weight:600; letter-spacing:2px; }}

/* рабочая область */
.content {{ position:absolute; left:96px; right:96px; top:168px; bottom:128px; z-index:2;
  display:flex; flex-direction:column; }}
.title {{ font-size:56px; line-height:1.08; font-weight:600; text-transform:uppercase; letter-spacing:.2px;
  max-width:1500px; }}
.title.s {{ font-size:48px; }} .title.xs {{ font-size:42px; }}
.bar {{ width:120px; height:6px; border-radius:3px; background:{GRAD_BAR}; margin:28px 0 40px; flex:none; }}
.dark .bar {{ background:#FFFFFF; }}
.lead {{ font-size:32px; line-height:1.45; color:{TEXT}; max-width:1400px; margin-bottom:40px; }}
.dark .lead {{ color:rgba(255,255,255,.92); }}
p.body {{ font-size:28px; line-height:1.5; color:{TEXT}; max-width:1400px; }}
.dark p.body {{ color:rgba(255,255,255,.92); }}
p + p {{ margin-top:18px; }}
p.body + ul.bl {{ margin-top:32px; }}
.muted {{ color:{MUTED}; }}
.label {{ font-size:18px; font-weight:700; letter-spacing:2.4px; text-transform:uppercase; color:{BLUE}; margin-bottom:14px; }}
.dark .label {{ color:#9FD8FF; }}
.source {{ margin-top:auto; padding-top:20px; font-size:18px; color:{MUTED}; }}
.dark .source {{ color:rgba(255,255,255,.65); }}

/* тезис-карточка */
.thesis {{ margin-top:auto; border-radius:24px; padding:34px 40px; background:{GRAD_CARD}; color:#FFFFFF;
  font-size:30px; line-height:1.4; font-weight:500; }}
.dark .thesis {{ background:rgba(255,255,255,.12); border:1px solid rgba(255,255,255,.22); }}
.thesis .label {{ color:rgba(255,255,255,.8); }}

/* обложка */
.cover .content {{ top:200px; }}
.cover h1 {{ font-size:96px; line-height:1.03; font-weight:500; text-transform:uppercase; letter-spacing:-.5px; max-width:1500px; }}
.cover h1.s {{ font-size:80px; }} .cover h1.xs {{ font-size:66px; }}
.cover .sub {{ margin-top:36px; font-size:34px; line-height:1.4; color:rgba(255,255,255,.92); max-width:1180px; }}
.chips {{ margin-top:auto; display:flex; flex-wrap:wrap; gap:14px; }}
.chip {{ padding:12px 24px; border-radius:999px; background:rgba(255,255,255,.14); border:1px solid rgba(255,255,255,.28);
  font-size:22px; line-height:1.3; color:#FFFFFF; white-space:nowrap; }}
.light .chip {{ background:{CARD}; border-color:{LINE}; color:{TEXT}; }}

/* раздел */
.section .num {{ font-size:260px; line-height:.85; font-weight:200; letter-spacing:-10px; margin:40px 0 48px; }}
.section h1 {{ font-size:84px; line-height:1.04; font-weight:500; text-transform:uppercase; max-width:1500px; }}
.section h1.s {{ font-size:68px; }}
.section .sub {{ margin-top:28px; font-size:32px; line-height:1.4; color:rgba(255,255,255,.88); max-width:1200px; }}

/* списки */
ul.bl {{ list-style:none; display:flex; flex-direction:column; gap:22px; max-width:1500px; }}
ul.bl li {{ position:relative; padding-left:44px; font-size:30px; line-height:1.42; }}
ul.bl li::before {{ content:''; position:absolute; left:0; top:15px; width:16px; height:16px; border-radius:4px;
  background:{GRAD_BAR}; transform:rotate(45deg) scale(.8); }}
.dark ul.bl li::before {{ background:#FFFFFF; }}
ul.bl li b {{ font-weight:700; }}
.cols2 {{ display:grid; grid-template-columns:1fr 1fr; gap:48px; flex:1; min-height:0; }}
.side {{ border-radius:24px; background:{CARD}; padding:36px 40px; font-size:26px; line-height:1.48; align-self:start; }}

/* повестка */
.agenda {{ display:grid; grid-template-columns:1fr 1fr; gap:22px 48px; counter-reset:a; }}
.agenda .it {{ display:flex; gap:28px; align-items:baseline; padding:22px 0; border-bottom:2px solid {LINE}; }}
.agenda .n {{ font-size:44px; font-weight:300; color:{BLUE}; min-width:72px; }}
.agenda .t {{ font-size:30px; line-height:1.3; font-weight:500; }}
.agenda .d {{ font-size:22px; color:{MUTED}; margin-top:6px; line-height:1.35; font-weight:400; }}

/* карточки */
.cards {{ display:grid; gap:24px; flex:1; min-height:0; align-content:start; }}
.card {{ border-radius:24px; background:{CARD}; padding:36px 36px 40px; display:flex; flex-direction:column; }}
.card .n {{ font-size:64px; line-height:1; font-weight:300; color:{BLUE}; margin-bottom:22px; }}
.card h3 {{ font-size:30px; line-height:1.22; font-weight:600; margin-bottom:14px; }}
.card p {{ font-size:24px; line-height:1.45; color:#3A404C; }}
.card.hl {{ background:{GRAD_CARD}; color:#FFFFFF; }}
.card.hl .n, .card.hl p {{ color:#FFFFFF; }}
.dark .card {{ background:rgba(255,255,255,.12); border:1px solid rgba(255,255,255,.22); }}
.dark .card .n {{ color:#9FD8FF; }} .dark .card p {{ color:rgba(255,255,255,.9); }}

/* цифры */
.stats {{ display:grid; gap:24px; }}
.stat {{ border-top:6px solid {BLUE}; padding-top:28px; }}
.stat .v {{ font-size:120px; line-height:.95; font-weight:300; letter-spacing:-3px; color:{BLUE}; white-space:nowrap; }}
.stat .v small {{ font-size:52px; letter-spacing:0; margin-left:6px; }}
.stat .l {{ margin-top:18px; font-size:26px; line-height:1.38; color:{TEXT}; }}
.dark .stat {{ border-color:#FFFFFF; }} .dark .stat .v {{ color:#FFFFFF; }} .dark .stat .l {{ color:rgba(255,255,255,.9); }}

/* процесс */
.process {{ display:grid; gap:16px; }}
.step {{ border-radius:20px; background:{CARD}; padding:30px 30px 34px; position:relative; }}
.step .k {{ display:flex; justify-content:space-between; align-items:center; font-size:22px; font-weight:700;
  letter-spacing:2px; color:{BLUE}; margin-bottom:22px; }}
.step .k i {{ font-style:normal; font-weight:400; font-size:30px; }}
.step h3 {{ font-size:28px; line-height:1.22; font-weight:600; margin-bottom:12px; }}
.step p {{ font-size:22px; line-height:1.42; color:#3A404C; }}
.step.last {{ background:{GRAD_CARD}; color:#FFFFFF; }}
.step.last .k, .step.last p {{ color:#FFFFFF; }}

/* сравнение */
.cmp {{ display:grid; grid-template-columns:1fr 1fr; gap:32px; align-items:start; }}
.cmp .col {{ border-radius:24px; padding:36px 40px; background:{CARD}; }}
.cmp .col.hl {{ background:{GRAD_CARD}; color:#FFFFFF; }}
.cmp h3 {{ font-size:34px; font-weight:600; text-transform:uppercase; margin-bottom:24px; }}
.cmp ul.bl li {{ font-size:26px; }}
.cmp .col.hl ul.bl li::before {{ background:#FFFFFF; }}

/* цитата */
.quote {{ position:relative; font-size:52px; line-height:1.28; font-weight:400; max-width:1500px; }}
.quote::before {{ content:'«'; position:absolute; left:-12px; top:-250px; font-size:240px; font-weight:200; opacity:.35; }}
.who {{ margin-top:48px; display:flex; align-items:center; gap:24px; }}
.who img {{ width:96px; height:96px; border-radius:50%; object-fit:cover; }}
.who .nm {{ font-size:28px; font-weight:600; }} .who .rl {{ font-size:22px; opacity:.8; margin-top:4px; }}

/* таблица */
table.t {{ width:100%; border-collapse:separate; border-spacing:0; border-radius:18px; overflow:hidden; font-size:24px; }}
table.t th {{ background:{BLUE}; color:#FFFFFF; font-weight:700; text-align:left; padding:20px 24px; font-size:22px; }}
table.t th:first-child {{ background:{NAVY}; }}
table.t td {{ padding:18px 24px; border-bottom:1px solid {LINE}; line-height:1.35; vertical-align:top; }}
table.t tr:nth-child(even) td {{ background:{CARD}; }}
table.t td.hl {{ color:{BLUE}; font-weight:700; }}
table.t td.num, table.t th.num {{ text-align:right; font-variant-numeric:tabular-nums; }}

/* график: горизонтальные столбцы */
.hb {{ display:flex; flex-direction:column; gap:26px; max-width:1500px; }}
.hb .row {{ display:grid; grid-template-columns:420px 1fr 160px; gap:28px; align-items:center; }}
.hb .nm {{ font-size:26px; line-height:1.25; }}
.hb .tr {{ height:30px; border-radius:15px; background:#E6EAF2; overflow:hidden; }}
.hb .fl {{ height:100%; border-radius:15px; background:{GRAD_BAR}; }}
.hb .fl.hl {{ background:linear-gradient(90deg,{VIOLET} 0%,{BLUE2} 100%); }}
.hb .v {{ font-size:32px; font-weight:600; color:{BLUE}; text-align:right; white-space:nowrap; }}

/* картинка + текст */
.media {{ display:grid; grid-template-columns:1fr 1fr; gap:64px; flex:1; min-height:0; }}
.media .ph {{ border-radius:28px; overflow:hidden; background:{CARD}; min-height:0; }}
.media .ph img {{ width:100%; height:100%; object-fit:cover; display:block; }}
.media .tx {{ display:flex; flex-direction:column; min-height:0; }}
.full-bleed {{ position:absolute; top:0; right:0; bottom:0; width:50%; }}
.full-bleed img {{ width:100%; height:100%; object-fit:cover; }}

/* человек (лектор / эксперт) */
.person {{ display:grid; grid-template-columns:520px 1fr; gap:72px; flex:1; min-height:0; }}
.person .ph {{ border-radius:28px; overflow:hidden; background:{GRAD_CARD}; height:100%; max-height:700px; }}
.person .ph img {{ width:100%; height:100%; object-fit:cover; display:block; }}
.person .nm {{ font-size:64px; line-height:1.05; font-weight:600; text-transform:uppercase; }}
.person .rl {{ margin-top:22px; font-size:30px; line-height:1.4; color:{BLUE}; font-weight:500; }}

/* финал */
.closing h1 {{ font-size:110px; line-height:1; font-weight:500; text-transform:uppercase; letter-spacing:-1px; }}
.closing .sub {{ margin-top:34px; font-size:34px; line-height:1.4; color:rgba(255,255,255,.92); max-width:1200px; }}
.contacts {{ margin-top:auto; display:flex; flex-wrap:wrap; gap:16px 56px; font-size:28px; }}
.contacts .c span {{ display:block; font-size:18px; letter-spacing:2px; text-transform:uppercase; opacity:.7; margin-bottom:8px; }}
.btn {{ display:inline-flex; align-items:center; gap:18px; margin-top:44px; padding:22px 40px; border-radius:16px;
  background:#FFFFFF; color:{VIOLET}; font-size:28px; font-weight:700; align-self:flex-start; }}
"""


RINGS = [
    [(1180, -360, 1000), (-420, 760, 760)],
    [(1100, 560, 1100), (-500, -420, 780)],
    [(1300, -500, 920), (-380, 820, 700)],
    [(960, 640, 1200), (1500, -260, 560)],
    [(-460, -420, 900), (1200, 620, 1000)],
]


def rings(i, light=False):
    if light:
        x, y, s = (1500, 720, 760)
        return f'<div class="ring" style="left:{x}px;top:{y}px;width:{s}px;height:{s}px"></div>'
    return "".join(f'<div class="ring" style="left:{x}px;top:{y}px;width:{s}px;height:{s}px"></div>'
                   for x, y, s in RINGS[i % len(RINGS)])


def size_cls(text, steps):
    """Класс уменьшения кегля по длине строки: steps = [(лимит, класс), …]."""
    n = len(re.sub(r"<[^>]+>|&nbsp;", " ", text or ""))
    for lim, cls in steps:
        if n > lim:
            return cls
    return ""


# --- Макеты ---------------------------------------------------------------
def title_block(s):
    t = typo(s.get("title", ""))
    cls = size_cls(s.get("title"), [(90, "xs"), (55, "s")])
    out = f'<h2 class="title {cls}">{t}</h2><div class="bar"></div>' if t else ""
    if s.get("lead"):
        out += f'<p class="lead">{typo(s["lead"])}</p>'
    return out


def paras(items):
    if isinstance(items, str):
        items = [items]
    return "".join(f'<p class="body">{typo(x)}</p>' for x in items or [])


def bullets_html(items):
    return '<ul class="bl">' + "".join(f"<li>{typo(x)}</li>" for x in items) + "</ul>"


def thesis_html(s):
    t = s.get("thesis")
    if not t:
        return ""
    lab = f'<div class="label">{typo(s["thesis_label"])}</div>' if s.get("thesis_label") else ""
    return f'<div class="thesis">{lab}{typo(t)}</div>'


def source_html(s):
    return f'<div class="source">Источник: {typo(s["source"])}</div>' if s.get("source") else ""


def gap(px):
    return f'<div style="height:{px}px;flex:none"></div>'


def L_cover(s, ctx):
    t = s.get("title", "")
    cls = size_cls(t, [(70, "xs"), (40, "s")])
    chips = "".join(f'<span class="chip">{typo(c)}</span>' for c in s.get("meta", []))
    sub = f'<p class="sub">{typo(s["subtitle"])}</p>' if s.get("subtitle") else ""
    return "dark cover", f'<h1 class="{cls}">{typo(t)}</h1>{sub}<div class="chips">{chips}</div>'


def L_section(s, ctx):
    num = s.get("num", "")
    cls = size_cls(s.get("title"), [(45, "s")])
    sub = f'<p class="sub">{typo(s["subtitle"])}</p>' if s.get("subtitle") else ""
    n = f'<div class="num">{esc(num)}</div>' if num != "" else gap(80)
    return "dark section", f'{n}<h1 class="{cls}">{typo(s.get("title", ""))}</h1>{sub}'


def L_agenda(s, ctx):
    items = []
    for i, it in enumerate(s["items"], start=1):
        if isinstance(it, str):
            it = {"title": it}
        d = f'<div class="d">{typo(it["text"])}</div>' if it.get("text") else ""
        items.append(f'<div class="it"><span class="n">{i:02d}</span><div><div class="t">{typo(it["title"])}</div>{d}</div></div>')
    return "light", title_block(s) + f'<div class="agenda">{"".join(items)}</div>'


def L_text(s, ctx):
    return "light", title_block(s) + paras(s.get("text")) + thesis_html(s) + source_html(s)


def L_bullets(s, ctx):
    lst = bullets_html(s["items"])
    if s.get("side"):
        side = s["side"]
        lab = f'<div class="label">{typo(side.get("label", ""))}</div>' if side.get("label") else ""
        body = f'<div class="cols2"><div>{lst}</div><div class="side">{lab}{typo(side.get("text", ""))}</div></div>'
    else:
        body = lst
    return "light", title_block(s) + body + thesis_html(s) + source_html(s)


def L_cards(s, ctx):
    items = s["items"]
    n = s.get("columns") or (len(items) if len(items) <= 4 else 3)
    cs = []
    for i, it in enumerate(items, start=1):
        num = it.get("num", f"{i:02d}") if s.get("numbered", True) else ""
        nh = f'<div class="n">{esc(num)}</div>' if num else ""
        hl = " hl" if it.get("highlight") else ""
        cs.append(f'<div class="card{hl}">{nh}<h3>{typo(it.get("title", ""))}</h3><p>{typo(it.get("text", ""))}</p></div>')
    theme = "dark" if s.get("dark") else "light"
    return theme, title_block(s) + f'<div class="cards" style="grid-template-columns:repeat({n},1fr)">{"".join(cs)}</div>' + thesis_html(s) + source_html(s)


def L_stats(s, ctx):
    items = s["items"]
    st = []
    for it in items:
        unit = f'<small>{esc(it["unit"])}</small>' if it.get("unit") else ""
        st.append(f'<div class="stat"><div class="v">{esc(it["value"])}{unit}</div><div class="l">{typo(it.get("label", ""))}</div></div>')
    theme = "dark" if s.get("dark") else "light"
    return theme, title_block(s) + gap(30) + f'<div class="stats" style="grid-template-columns:repeat({len(items)},1fr)">{"".join(st)}</div>' + thesis_html(s) + source_html(s)


def L_process(s, ctx):
    items = s["items"]
    st = []
    for i, it in enumerate(items, start=1):
        last = i == len(items)
        arrow = "" if last else "<i>→</i>"
        st.append(f'<div class="step{" last" if last and s.get("highlight_last", True) else ""}"><div class="k"><span>ШАГ {i}</span>{arrow}</div>'
                  f'<h3>{typo(it.get("title", ""))}</h3><p>{typo(it.get("text", ""))}</p></div>')
    return "light", title_block(s) + f'<div class="process" style="grid-template-columns:repeat({len(items)},1fr)">{"".join(st)}</div>' + thesis_html(s) + source_html(s)


def L_compare(s, ctx):
    cols = []
    for c in s["columns"][:2]:
        hl = " hl" if c.get("highlight") else ""
        cols.append(f'<div class="col{hl}"><h3>{typo(c["title"])}</h3>{bullets_html(c["items"])}</div>')
    return "light", title_block(s) + f'<div class="cmp">{"".join(cols)}</div>' + thesis_html(s)


def L_quote(s, ctx):
    who = ""
    if s.get("author"):
        img = f'<img src="{data_uri(ctx.res(s["photo"]))}">' if s.get("photo") else ""
        rl = f'<div class="rl">{typo(s["role"])}</div>' if s.get("role") else ""
        who = f'<div class="who">{img}<div><div class="nm">{typo(s["author"])}</div>{rl}</div></div>'
    return "dark", f'{gap(150)}<div class="quote">{typo(s["text"])}</div>{who}'


def L_table(s, ctx):
    hdr = s["headers"]
    align = s.get("align", [])
    hl = {tuple(x) for x in s.get("highlight", [])}
    def cls(i, extra=""):
        c = ("num " if i < len(align) and align[i] == "right" else "") + extra
        return f' class="{c.strip()}"' if c.strip() else ""
    th = "".join(f"<th{cls(i)}>{typo(h)}</th>" for i, h in enumerate(hdr))
    trs = "".join("<tr>" + "".join(f'<td{cls(j, "hl" if (r, j) in hl else "")}>{typo(v)}</td>' for j, v in enumerate(row)) + "</tr>"
                  for r, row in enumerate(s["rows"]))
    return "light", title_block(s) + f'<table class="t"><thead><tr>{th}</tr></thead><tbody>{trs}</tbody></table>' + source_html(s)


def L_chart(s, ctx):
    items = s["items"]
    vmax = s.get("max") or max(float(i["value"]) for i in items)
    unit = s.get("unit", "")
    rows = []
    for it in items:
        w = max(0.0, min(100.0, float(it["value"]) / vmax * 100))
        v = it.get("display", f'{it["value"]:g}' if isinstance(it["value"], (int, float)) else it["value"])
        rows.append(f'<div class="row"><div class="nm">{typo(it["label"])}</div><div class="tr"><div class="fl{" hl" if it.get("highlight") else ""}" style="width:{w:.1f}%"></div></div><div class="v">{esc(v)}{esc(unit)}</div></div>')
    return "light", title_block(s) + f'<div class="hb">{"".join(rows)}</div>' + thesis_html(s) + source_html(s)


def L_image(s, ctx):
    img = f'<img src="{data_uri(ctx.res(s["image"]))}">' if s.get("image") else ""
    tx = title_block(s) + paras(s.get("text")) + (bullets_html(s["items"]) if s.get("items") else "") + thesis_html(s)
    left, right = (f'<div class="ph">{img}</div>', f'<div class="tx">{tx}</div>')
    if s.get("image_side", "right") == "right":
        left, right = right, left
    return "light", f'<div class="media">{left}{right}</div>'


def L_person(s, ctx):
    img = f'<img src="{data_uri(ctx.res(s["photo"]))}">' if s.get("photo") else ""
    pts = bullets_html(s["items"]) if s.get("items") else ""
    rl = f'<div class="rl">{typo(s["role"])}</div>' if s.get("role") else ""
    return "light", (f'<div class="person"><div class="ph">{img}</div><div style="display:flex;flex-direction:column">'
                     f'<div class="label">{typo(s.get("label", "Лектор"))}</div><div class="nm">{typo(s["name"])}</div>{rl}'
                     f'<div class="bar"></div>{paras(s.get("text"))}{pts}</div></div>')


def L_closing(s, ctx):
    sub = f'<p class="sub">{typo(s["subtitle"])}</p>' if s.get("subtitle") else ""
    btn = f'<span class="btn">{typo(s["button"])} <b>→</b></span>' if s.get("button") else ""
    cs = "".join(f'<div class="c"><span>{typo(c["label"])}</span>{typo(c["value"])}</div>' for c in s.get("contacts", []))
    return "dark closing", f'{gap(60)}<h1>{typo(s.get("title", "Спасибо!"))}</h1>{sub}{btn}<div class="contacts">{cs}</div>'


LAYOUTS = {k[2:]: v for k, v in globals().items() if k.startswith("L_")}


class Ctx:
    def __init__(self, base):
        self.base = base

    def res(self, p):
        q = Path(p)
        return q if q.is_absolute() else (self.base / q)


def build_html(deck, base):
    ctx = Ctx(base)
    font = data_uri(ASSETS / "fonts" / "InterVariable.ttf")
    logo_w = data_uri(ASSETS / "logo" / "vtb-edu-logo-white.png")
    logo_d = data_uri(ASSETS / "logo" / "vtb-edu-logo-navy.png")
    slides = deck["slides"]
    total = len(slides)
    footer = deck.get("footer", "")
    kicker_default = deck.get("kicker", "")
    out = []
    for i, s in enumerate(slides, start=1):
        lay = s.get("layout", "text")
        if lay not in LAYOUTS:
            sys.exit(f"Слайд {i}: неизвестный макет «{lay}». Доступны: {', '.join(LAYOUTS)}")
        theme, inner = LAYOUTS[lay](s, ctx)
        dark = theme.startswith("dark")
        logo = f'<img class="logo" src="{logo_w if dark else logo_d}" alt="ВТБ Образование">'
        k = s.get("kicker", kicker_default)
        kick = f'<div class="kicker">{typo(k)}</div>' if k else ""
        pg = "" if lay == "cover" or s.get("hide_page") else f'<span class="pg">{i:02d} / {total:02d}</span>'
        foot = f'<div class="foot"><span>{typo(s.get("footer", footer))}</span>{pg}</div>'
        bg = rings(i, light=not dark) if (dark or s.get("rings")) else ""
        notes = esc(s.get("notes", ""))
        out.append(f'<section class="slide {theme}" data-layout="{lay}" data-notes="{notes}">{bg}{logo}{kick}'
                   f'<div class="content">{inner}</div>{foot}</section>')
    title = esc(deck.get("title", "Презентация"))
    return (f'<!doctype html><html lang="ru"><head><meta charset="utf-8"><title>{title}</title>'
            f'<style>{css(font)}</style></head><body>{"".join(out)}</body></html>')


# --- Рендер ---------------------------------------------------------------
CHECK_JS = """() => [...document.querySelectorAll('.slide')].map((sl, i) => {
  const c = sl.querySelector('.content'), cb = c.getBoundingClientRect();
  const bad = [...c.querySelectorAll('*')].filter(e => { const r = e.getBoundingClientRect();
    return r.width && (r.bottom > cb.bottom + 1 || r.right > cb.right + 1); }).length;
  const scroll = c.scrollHeight > c.clientHeight + 1;
  return {slide: i + 1, layout: sl.dataset.layout, overflow: bad > 0 || scroll};
})"""


def render(html_path, out):
    try:
        from playwright.sync_api import sync_playwright  # noqa
        return render_py(html_path, out)
    except ImportError:
        return render_node(html_path, out)


def render_py(html_path, out):
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        try:
            b = p.chromium.launch()
        except Exception:
            b = p.chromium.launch(executable_path="/opt/pw-browsers/chromium")
        pg = b.new_page(viewport={"width": W, "height": H})
        pg.goto(html_path.as_uri())
        pg.evaluate("document.fonts.ready")
        report = pg.evaluate(CHECK_JS)
        for i, el in enumerate(pg.query_selector_all(".slide"), start=1):
            el.screenshot(path=str(out / f"slide_{i:02d}.png"))
        pg.pdf(path=str(out / "deck.pdf"), width=f"{W}px", height=f"{H}px", print_background=True)
        b.close()
    return report


def render_node(html_path, out):
    if not shutil.which("node"):
        sys.exit("Нужен Playwright: pip install playwright && playwright install chromium (или node + playwright).")
    js = Path(__file__).with_name("render.cjs")
    r = subprocess.run(["node", str(js), str(html_path), str(out), str(W), str(H), CHECK_JS],
                       capture_output=True, text=True, cwd="/")
    if r.returncode:
        sys.exit(r.stderr)
    return json.loads(r.stdout)


def build_pptx(deck, out, n):
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        print("python-pptx не установлен — deck.pptx пропущен (pip install python-pptx)")
        return
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    blank = prs.slide_layouts[6]
    for i in range(1, n + 1):
        sl = prs.slides.add_slide(blank)
        sl.shapes.add_picture(str(out / f"slide_{i:02d}.png"), 0, 0, prs.slide_width, prs.slide_height)
        note = deck["slides"][i - 1].get("notes")
        if note:
            sl.notes_slide.notes_text_frame.text = note
    prs.core_properties.title = deck.get("title", "")
    prs.save(out / "deck.pptx")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck")
    ap.add_argument("-o", "--out", default="out")
    ap.add_argument("--html-only", action="store_true")
    a = ap.parse_args()
    src = Path(a.deck).resolve()
    deck = json.loads(src.read_text(encoding="utf-8"))
    out = Path(a.out).resolve()
    out.mkdir(parents=True, exist_ok=True)
    doc = build_html(deck, src.parent)
    hp = out / "deck.html"
    hp.write_text(doc.replace("<body>", '<body class="deck-preview" style="padding:40px 0">', 1), encoding="utf-8")
    if a.html_only:
        print(hp)
        return
    tmp = out / ".render.html"
    tmp.write_text(doc, encoding="utf-8")
    report = render(tmp, out)
    tmp.unlink()
    build_pptx(deck, out, len(deck["slides"]))
    (out / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=1), encoding="utf-8")
    bad = [r for r in report if r["overflow"]]
    for r in report:
        print(f'slide_{r["slide"]:02d}.png  {r["layout"]:<8} {"ПЕРЕПОЛНЕНИЕ" if r["overflow"] else "ok"}')
    print(f"\nГотово: {out}  (слайдов: {len(report)}, переполнено: {len(bad)})")
    if bad:
        sys.exit(2)


if __name__ == "__main__":
    main()
