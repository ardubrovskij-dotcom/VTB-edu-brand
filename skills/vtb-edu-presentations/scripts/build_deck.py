#!/usr/bin/env python3
"""deck.json → редактируемый PPTX по шаблону «мерико × ВТБ Образование».

Каждый слайд — копия слайда-прототипа из assets/template.pptx, привязанная к мастеру нужного направления.
Скрипт заполняет плейсхолдеры, меняет иконки, фото, портреты, таблицы и диаграммы, убирает лишние элементы.

    python3 build_deck.py deck.json -o out/ --wireframe   # 1) прототип: серые блоки для согласования
    python3 build_deck.py deck.json -o out/ --preview     # 2) дизайн после согласования + ревью
    python3 build_deck.py --list-icons
"""
import argparse, copy, io, json, os, re, shutil, subprocess, sys
from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import PartFactory
from pptx.opc.packuri import PackURI
from pptx.oxml import parse_xml
from pptx.dml.color import RGBColor
from pptx.util import Pt, Emu

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from layouts import LAYOUTS, DIRECTIONS, PEOPLE  # noqa: E402
import house  # noqa: E402

if __name__ == '__main__':
    # house.py обращается к build_deck через import: без этого он получит второй экземпляр модуля со своим
    # UNIFORM, и единый кегль для макетов фирменного стиля не сработает
    sys.modules.setdefault('build_deck', sys.modules[__name__])

SKILL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ASSETS = os.path.join(SKILL, 'assets')
E = 12700  # EMU в 1 pt
NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
R_NS = NS['r']
MIN_PT = 9
# Шрифт бренда с 2026 г. — VTB Omega UI (Regular, Medium, DemiBold, Bold). VTB Group шаблона заменяется при сборке.
FONT = 'VTB Omega UI'
FONT_DEMI = 'VTB Omega UI Medium'  # выделение — Medium (DemiBold слишком жирный)
# Единый кегль по колоде: если элемент пришлось уменьшить на одном слайде, тот же элемент (роль) на всех
# слайдах получает тот же кегль. shrunk — минимальный кегль роли, до которого пришлось уменьшать в этом проходе,
# caps — потолок кегля роли во втором проходе.
UNIFORM = {'caps': {}, 'shrunk': {}}
COVERS = {'cover', 'cover-panel', 'cover-band', 'cover-text', 'cover-diagonal', 'cover-photo', 'section', 'title', 'chapter'}


def role_of(where):
    m = re.search(r'\(([^)]+)\)/([^/]+)$', where)
    if not m:
        return None
    raw = m.group(2)
    field = re.sub(r'\d+$', '', raw)
    if field != raw:
        field = 'элемент:' + field  # title1, text2 — поля карточек, а не заголовок слайда
    return ('обложка:' if m.group(1) in COVERS else '') + field


def capped(role, size):
    return min(size, UNIFORM['caps'].get(role, size)) if role else size


def note_size(role, base, final):
    if role and final < base - 0.01:
        UNIFORM['shrunk'][role] = min(final, UNIFORM['shrunk'].get(role, final))


FONT_MAP = {'VTB Group Book': FONT, 'VTB Group Light': FONT, 'VTB Group': FONT,
            'VTB Group Bold': FONT_DEMI, 'VTB Group Medium': FONT_DEMI, 'VTB Group DemiBold': FONT_DEMI}
LIGHT_DIRECTIONS = {5, 6}  # b2b, infra: светлые фоны и карточки
DARK_TEXT = '0E2569'  # тёмный текст мастера «Инфраструктурные проекты»


def stype(sh):
    try:
        return sh.shape_type
    except NotImplementedError:
        return None


def xp(el, q):
    return etree._Element.xpath(el, q, namespaces=NS)

# ---------------------------------------------------------------- библиотека иконок и иллюстраций
# Слайды 63–69 шаблона: верхний ряд — «Градиент», нижний — «Заливка». Имена по порядку слева направо.
ICON_ROWS = {
    63: (['sign', 'edit', 'graduation', 'search', 'timer', 'bulb', 'presentation'],) * 2,
    64: (['edit', 'briefcase', 'calculator', 'money', 'bank', 'card'],) * 2,
    65: (['globe', 'folder-in', 'link', 'server', 'monitor', 'cloud', 'coins'],) * 2,
    66: (['people', 'people-list', 'user-add', 'user-target', 'user-growth', 'user-chat'],
         ['people', 'people-list', 'user-add', 'user-growth', 'user-target', 'user-chat']),
    67: (['sign', 'edit', 'image', 'layers', 'mail', 'send', 'folder', 'structure', 'dots'],
         ['sign', 'edit', 'image', 'layers', 'mail', 'send', 'folder', 'structure', 'list']),
    68: (['globe', 'people', 'mail', 'send', 'phone', 'chat', 'handshake'],) * 2,
    69: (['badge', 'bell', 'pin', 'home', 'copy', 'download', 'star', 'gear'],) * 2,
}
ILL_SLIDES = {70: 'edu', 71: 'finance', 72: 'people', 73: 'management', 74: 'it', 75: 'other'}


def bbox(el_shape):
    return (el_shape.left / E, el_shape.top / E, el_shape.width / E, el_shape.height / E)


def clusters(shapes, pad=6):
    """Группирует фигуры, чьи рамки пересекаются: одна иконка может состоять из нескольких фигур."""
    items = [[s] for s in shapes]
    boxes = [bbox(s) for s in shapes]
    changed = True
    while changed:
        changed = False
        for i in range(len(items)):
            for j in range(i + 1, len(items)):
                a, b = boxes[i], boxes[j]
                if a[0] - pad < b[0] + b[2] and b[0] - pad < a[0] + a[2] and a[1] - pad < b[1] + b[3] and b[1] - pad < a[1] + a[3]:
                    items[i] += items.pop(j)
                    x0, y0 = min(a[0], b[0]), min(a[1], b[1])
                    x1, y1 = max(a[0] + a[2], b[0] + b[2]), max(a[1] + a[3], b[1] + b[3])
                    boxes[i] = (x0, y0, x1 - x0, y1 - y0)
                    boxes.pop(j)
                    changed = True
                    break
            if changed:
                break
    return list(zip(items, boxes))


def rows_of(cl, tol=40):
    cl = sorted(cl, key=lambda c: c[1][1] + c[1][3] / 2)
    rows = []
    for c in cl:
        cy = c[1][1] + c[1][3] / 2
        if rows and abs(rows[-1][0] - cy) < tol:
            rows[-1][1].append(c)
        else:
            rows.append([cy, [c]])
    return [sorted(r[1], key=lambda c: c[1][0]) for r in rows]


def build_library(prs):
    lib = {}
    for s in prs.slides:
        m = re.match(r'proto-(\d+)', s._element.cSld.get('name', ''))
        if not m:
            continue
        n = int(m.group(1))
        if n not in ICON_ROWS and n not in ILL_SLIDES:
            continue
        shapes = [sh for sh in s.shapes if not sh.is_placeholder and stype(sh) in (1, 5, 6, 13)
                  and sh.top / E > 95 and not (sh.has_text_frame and sh.text_frame.text.strip())]
        rows = rows_of(clusters(shapes))
        if n in ICON_ROWS:
            for style, names, row in zip(('', '-fill'), ICON_ROWS[n], rows):
                for name, c in zip(names, row):
                    lib.setdefault(name + style, (s, c))
        else:
            k = 0
            for row in rows:
                for c in row:
                    k += 1
                    lib[f'ill-{ILL_SLIDES[n]}-{k:02d}'] = (s, c)
    return lib


# ---------------------------------------------------------------- копирование частей пакета
def copy_part(part, package, cache):
    if id(part) in cache:
        return cache[id(part)]
    name = str(part.partname).replace('%', '')
    tmpl = re.sub(r'\d*(\.\w+)$', r'%d\1', name)
    new = PartFactory(package.next_partname(tmpl), part.content_type, package, part.blob)
    cache[id(part)] = new
    remap = {}
    for rId, rel in sorted(part.rels.items(), key=lambda kv: int(re.sub(r'\D', '', kv[0]) or 0)):
        if rel.is_external:
            remap[rId] = new.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            remap[rId] = new.relate_to(copy_part(rel.target_part, package, cache), rel.reltype)
    if hasattr(new, '_element') and any(k != v for k, v in remap.items()):
        remap_rids(new._element, remap)
    return new


def remap_rids(el, remap):
    for node in el.iter():
        for k, v in list(node.attrib.items()):
            if k.startswith('{%s}' % R_NS) and v in remap:
                node.set(k, remap[v])


def relate_from(src_part, dst_part, rIds, package, deep_types=(RT.CHART,)):
    remap, cache = {}, {}
    for rId in rIds:
        rel = src_part.rels.get(rId) if hasattr(src_part.rels, 'get') else src_part.rels[rId]
        if rel is None:
            continue
        if rel.is_external:
            remap[rId] = dst_part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        elif rel.reltype in deep_types:
            remap[rId] = dst_part.relate_to(copy_part(rel.target_part, package, cache), rel.reltype)
        else:
            remap[rId] = dst_part.relate_to(rel.target_part, rel.reltype)
    return remap


def rids_in(el):
    return {v for node in el.iter() for k, v in node.attrib.items() if k.startswith('{%s}' % R_NS)}


def clone_slide(prs, proto, layout):
    new = prs.slides.add_slide(layout)
    csld = copy.deepcopy(proto._element.cSld)
    remap = relate_from(proto.part, new.part, rids_in(csld), prs.part.package)
    remap_rids(csld, remap)
    # меняем содержимое на месте: python-pptx кэширует объект spTree у слайда
    dst = new._element.cSld
    tree = dst.spTree
    for ch in list(dst):
        if ch is not tree:
            dst.remove(ch)
    for ch in list(tree):
        tree.remove(ch)
    for ch in list(csld):
        if ch.tag == tree.tag:
            for sub in list(ch):
                tree.append(sub)
        else:
            tree.addprevious(ch) if ch.tag.endswith('}bg') else dst.append(ch)
    return new


# ---------------------------------------------------------------- утилиты фигур
def center_in(sh, box):
    x, y, w, h = bbox(sh)
    cx, cy = x + w / 2, y + h / 2
    return box[0] <= cx <= box[0] + box[2] and box[1] <= cy <= box[1] + box[3]


def delete(sh):
    el = sh._element
    el.getparent().remove(el)


def ph_by_idx(slide, idx):
    for sh in slide.placeholders:
        if sh.placeholder_format.idx == idx:
            return sh
    return None


def find_shape(slide, prefix):
    for sh in slide.shapes:
        if sh.name.startswith(prefix):
            return sh
    return None


def asset(kind, name, exts=('.png', '.jpg', '.jpeg')):
    if not name:
        return None
    if os.path.isfile(name):
        return name
    for e in ('',) + tuple(exts):
        p = os.path.join(ASSETS, kind, name + e)
        if os.path.isfile(p):
            return p
    return None


def person_key(name):
    """Портрет: ключ (artyom), русское имя (Артём) или направление (prof) → имя файла в assets/people."""
    if not isinstance(name, str):
        return name
    low = name.strip().lower().replace('ё', 'е')
    for key, (ru, _, direction) in PEOPLE.items():
        if low in (key, ru.lower().replace('ё', 'е'), direction):
            return key
    return name


# ---------------------------------------------------------------- текст и подбор кегля
class Fitter:
    def __init__(self):
        self.font_path = self._find_font()
        self.cache = {}

    @staticmethod
    def _find_font():
        try:
            out = subprocess.run(['fc-match', '-f', '%{file}', FONT], capture_output=True, text=True).stdout
            if out and 'omega' in out.lower():
                return out
        except FileNotFoundError:
            pass
        return os.path.join(ASSETS, 'fonts', 'OpenSans-Regular.ttf')

    def _demi(self):
        if not hasattr(self, '_demi_path'):
            self._demi_path = None
            try:
                out = subprocess.run(['fc-match', '-f', '%{file}', FONT_DEMI], capture_output=True, text=True).stdout
                if out and 'medium' in out.lower():
                    self._demi_path = out
            except FileNotFoundError:
                pass
        return self._demi_path

    def width(self, text, size, bold=False):
        path = (self._demi() if bold else None) or self.font_path
        key = (path, round(size * 4))
        f = self.cache.get(key) or self.cache.setdefault(key, ImageFont.truetype(path, key[1] * 10))
        w = f.getlength(text) / 40 * 1.04  # запас: PowerPoint набирает чуть шире, чем замер Pillow
        return w * (1.06 if bold and path == self.font_path else 1)

    def lines(self, text, size, width, caps=False, bold=False):
        n = 0
        for para in text.split('\n'):
            para = para.upper() if caps else para
            words, cur = para.split(' '), ''
            n += 1
            for w in words:
                t = (cur + ' ' + w).strip()
                if self.width(t, size, bold) <= width or not cur:
                    cur = t
                else:
                    n += 1
                    cur = w
        return n


def resolve(el, xpath):
    r = xp(el, xpath)
    return r[0] if r else None


def eff_size(sh, slide):
    """Кегль плейсхолдера: прямой → layout → master."""
    sizes = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
    if sizes:
        return max(sizes)
    chain = []
    if sh.is_placeholder:
        idx = sh.placeholder_format.idx
        lay = slide.slide_layout
        lp = next((p for p in lay.placeholders if p.placeholder_format.idx == idx), None)
        chain.append(sh._element)
        if lp is not None:
            chain.append(lp._element)
            for r in lp.text_frame.paragraphs:
                for run in r.runs:
                    if run.font.size:
                        return run.font.size.pt
        typ = 'title' if sh.placeholder_format.type in (1, 3) else 'body'
        m = lay.slide_master
        chain.append(resolve(m._element, './/p:txStyles/p:%sStyle' % ('title' if typ == 'title' else 'body')))
    for el in chain:
        if el is None:
            continue
        v = resolve(el, './/a:lvl1pPr/a:defRPr/@sz')
        if v:
            return int(v) / 100
    return 14.0


def insets(sh, slide):
    els = [sh._element]
    if sh.is_placeholder:
        els += [p._element for p in slide.slide_layout.placeholders
                if p.placeholder_format.idx == sh.placeholder_format.idx]
    total = 0
    for attr in ('lIns', 'rIns'):
        v = None
        for el in els:
            r = xp(el, './/a:bodyPr/@' + attr)
            if r:
                v = int(r[0]) / E
                break
        total += 7.2 if v is None else v
    return total


def is_caps(sh, slide):
    for el in [sh._element] + [p._element for p in slide.slide_layout.placeholders
                               if sh.is_placeholder and p.placeholder_format.idx == sh.placeholder_format.idx]:
        if xp(el, './/a:defRPr[@cap="all"]|.//a:rPr[@cap="all"]'):
            return True
    return False


def set_text(sh, value, slide, fitter, report, where, max_lines=None):
    if isinstance(value, (int, float)):
        value = str(value)
    if isinstance(value, list):
        value = '\n'.join(map(str, value))
    value = value.replace('\\n', '\n')
    tf = sh.text_frame
    # сохраняем формат первого абзаца и первого рана прототипа
    p0 = tf.paragraphs[0]._p
    pPr = copy.deepcopy(p0.find('a:pPr', NS))
    r0 = p0.find('a:r', NS)
    rPr = copy.deepcopy(r0.find('a:rPr', NS)) if r0 is not None and r0.find('a:rPr', NS) is not None else None
    txBody = tf._txBody
    bp = txBody.find('a:bodyPr', NS)
    for af in bp.findall('a:normAutofit', NS) if bp is not None else []:
        bp.remove(af)  # сжатие PowerPoint из прототипа не нужно: кегль подбирает fit()
    for p in txBody.findall('a:p', NS):
        txBody.remove(p)
    for line in value.split('\n'):
        p = etree.SubElement(txBody, '{%s}p' % NS['a'])
        if pPr is not None:
            p.append(copy.deepcopy(pPr))
        r = etree.SubElement(p, '{%s}r' % NS['a'])
        if rPr is not None:
            r.append(copy.deepcopy(rPr))
        else:
            etree.SubElement(r, '{%s}rPr' % NS['a'], lang='ru-RU')
        t = etree.SubElement(r, '{%s}t' % NS['a'])
        t.text = nbsp(line)
    fit(sh, slide, fitter, report, where, max_lines)


# слова, которые не оставляют в конце строки: предлоги, союзы, частицы, числа (перед единицей)
HANGING = r'[а-яА-ЯёЁa-zA-Z]{1,2}|\d+|(?:[Дд]ля|[Пп]ри|[Бб]ез|[Пп]ро|[Нн]ад|[Пп]од|[Оо]бо|[Ии]ли|[Чч]то|[Кк]ак|[Нн]е|[Нн]о)'


def nbsp(s):
    """Неразрывный пробел после предлогов, союзов и чисел и перед тире — без висячих слов в конце строки."""
    s = re.sub(r'(?<![\w\-])(' + HANGING + r') (?=\S)', lambda m: m.group(1) + ' ', s)
    s = re.sub(r' —', ' —', s)
    return s


def fit(sh, slide, fitter, report, where, max_lines=None):
    size = eff_size(sh, slide)
    w = sh.width / E - insets(sh, slide)
    h = sh.height / E
    caps = is_caps(sh, slide)
    text = sh.text_frame.text
    if not text.strip() or w <= 0:
        return
    role = role_of(where)
    s = capped(role, size)
    while True:
        n = fitter.lines(text, s, w, caps)
        # самое длинное слово тоже должно влезать: иначе PowerPoint разорвёт его по буквам («0 / 1»)
        widest = max((fitter.width(word.upper() if caps else word, s) for word in text.split()), default=0)
        if widest > w:
            n = max(n, (max_lines or 1) + 1)
        lim = max_lines or max(1, int((h * 1.25) // (s * 1.2)))
        if n <= lim or s <= max(MIN_PT, size * (0.5 if max_lines == 1 else 0.7)):
            break
        s -= 0.5
    if max_lines == 1:  # однострочное поле: без переноса (номер, год, цифра)
        bp = sh.text_frame._txBody.find('a:bodyPr', NS)
        if bp is not None:
            bp.set('wrap', 'none')
    note_size(role, size, s)
    if s < size:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(s)
        report.append(f'{where}: кегль {size:g} → {s:g} pt')
    if n > lim:
        report.append(f'{where}: ПЕРЕПОЛНЕНИЕ — {n} строк при пределе {lim}. Сократите текст: «{text[:60]}…»')


# ---------------------------------------------------------------- картинки
def put_image(slide, pic, path, focus=(0.5, 0.4)):
    """Заменяет картинку, сохраняя рамку: кадрирует «по заполнению» с фокусом."""
    img_part, rId = slide.part.get_or_add_image_part(path)
    blip = xp(pic._element, './/a:blip')[0]
    blip.set('{%s}embed' % R_NS, rId)
    for ext in blip.findall('a:extLst', NS):  # svg-версия больше не нужна
        blip.remove(ext)
    with Image.open(path) as im:
        iw, ih = im.size
    fw, fh = pic.width, pic.height
    ia, fa = iw / ih, fw / fh
    l = r = t = b = 0.0
    if ia > fa:
        cut = 1 - fa / ia
        l = cut * focus[0]
        r = cut - l
    else:
        cut = 1 - ia / fa
        t = cut * focus[1]
        b = cut - t
    pic.crop_left, pic.crop_right, pic.crop_top, pic.crop_bottom = l, r, t, b


def swap_blob(slide, pic, path):
    img_part, rId = slide.part.get_or_add_image_part(path)
    blip = xp(pic._element, './/a:blip')[0]
    blip.set('{%s}embed' % R_NS, rId)


def place_icon(slide, lib, name, target, prs):
    """Ставит иконку из библиотеки шаблона в рамку target (x, y, w, h в pt), сохраняя пропорции."""
    if name not in lib:
        return False
    src_slide, (shapes, (bx, by, bw, bh)) = lib[name]
    tx, ty, tw, th = target
    k = min(tw / bw, th / bh)
    nw, nh = bw * k, bh * k
    nx, ny = tx + (tw - nw) / 2, ty + (th - nh) / 2
    grp = etree.SubElement(slide.shapes._spTree, '{%s}grpSp' % NS['p'])
    nv = etree.SubElement(grp, '{%s}nvGrpSpPr' % NS['p'])
    etree.SubElement(nv, '{%s}cNvPr' % NS['p'], id=str(next_id(slide)), name=f'Иконка {name}')
    etree.SubElement(nv, '{%s}cNvGrpSpPr' % NS['p'])
    etree.SubElement(nv, '{%s}nvPr' % NS['p'])
    gp = etree.SubElement(grp, '{%s}grpSpPr' % NS['p'])
    xf = etree.SubElement(gp, '{%s}xfrm' % NS['a'])
    for tag, a, b in (('off', nx, ny), ('ext', nw, nh), ('chOff', bx, by), ('chExt', bw, bh)):
        etree.SubElement(xf, '{%s}%s' % (NS['a'], tag), **({'x': str(int(a * E)), 'y': str(int(b * E))}
                                                             if tag in ('off', 'chOff') else
                                                             {'cx': str(int(a * E)), 'cy': str(int(b * E))}))
    for sh in shapes:
        el = copy.deepcopy(sh._element)
        remap = relate_from(src_slide.part, slide.part, rids_in(el), prs.part.package)
        remap_rids(el, remap)
        # толщина обводки в группе не масштабируется: без этого иконка 16 pt рисуется линиями от 40 pt
        # и превращается в кляксу (structure, money, timer)
        for ln in el.iter('{%s}ln' % NS['a']):
            if ln.get('w'):
                ln.set('w', str(max(3175, int(int(ln.get('w')) * k))))
        grp.append(el)
    renumber(slide)
    return True


def next_id(slide):
    ids = [int(x) for x in xp(slide.shapes._spTree, './/p:cNvPr/@id')]
    return max(ids + [1]) + 1


def renumber(slide):
    seen, nid = set(), next_id(slide)
    for c in xp(slide.shapes._spTree, './/p:cNvPr'):
        if c.get('id') in seen:
            c.set('id', str(nid))
            nid += 1
        seen.add(c.get('id'))


# ---------------------------------------------------------------- таблицы и диаграммы
def fill_table(gf, data, report, where, fitter=None):
    tbl = xp(gf._element, './/a:tbl')[0]
    cols = data.get('columns') or []
    rows = data.get('rows') or []
    ncol = max([len(cols)] + [len(r) for r in rows]) or 1
    grid = tbl.find('a:tblGrid', NS)
    gcols = grid.findall('a:gridCol', NS)
    trs = tbl.findall('a:tr', NS)
    head, body = trs[0], trs[1:]
    total_w = sum(int(g.get('w')) for g in gcols)
    total_h = gf.height

    def fit_cols(tr, n):
        tcs = tr.findall('a:tc', NS)
        while len(tcs) < n:
            src = tcs[1] if len(tcs) > 1 else tcs[-1]
            new = copy.deepcopy(src)
            tr.insert(list(tr).index(tcs[-1]) + 1, new)
            tcs.append(new)
        for tc in tcs[n:]:
            tr.remove(tc)
    while len(gcols) < ncol:
        g = copy.deepcopy(gcols[-1]); grid.append(g); gcols.append(g)
    for g in gcols[ncol:]:
        grid.remove(g)
    gcols = gcols[:ncol]
    # ширины колонок — по длине текста, но не уже 14 % ширины таблицы
    lens = [max([len(str(c)) for c in [cols[i] if i < len(cols) else ''] + [r[i] if i < len(r) else '' for r in rows]] + [4])
            for i in range(ncol)]
    weights = [max(0.14, l / sum(lens)) for l in lens]
    weights = [w / sum(weights) for w in weights]
    for i, g in enumerate(gcols):
        g.set('w', str(int(total_w * weights[i])))
        for ext in g.findall('a:extLst', NS):
            g.remove(ext)
    # строки: чередование фона как в прототипе (первые две строки тела)
    patterns = body[:2] or [head]
    for tr in body:
        tbl.remove(tr)
    nrows = len(rows)
    row_h = int(min(total_h / (nrows + 1), 44 * E))
    for tr in [head]:
        fit_cols(tr, ncol)
        for ext in tr.findall('a:extLst', NS):
            tr.remove(ext)
        tr.set('h', str(row_h))
        for tc, v in zip(tr.findall('a:tc', NS), cols + [''] * ncol):
            cell_text(tc, v)
    for i, r in enumerate(rows):
        tr = copy.deepcopy(patterns[i % len(patterns)])
        for ext in tr.findall('a:extLst', NS):
            tr.remove(ext)
        tr.set('h', str(row_h))
        fit_cols(tr, ncol)
        for tc, v in zip(tr.findall('a:tc', NS), list(r) + [''] * ncol):
            cell_text(tc, v)
        tbl.append(tr)
    if fitter is not None:
        fit_table(tbl, gf.height / E, fitter, report, where)
    if nrows > 12:
        report.append(f'{where}: в таблице {nrows} строк — больше 12 плохо читается, разбейте на два слайда')


def fit_table(tbl, avail, fitter, report, where, start=16, floor=11):
    """Подбирает один кегль для всей таблицы так, чтобы строки с переносами уместились в высоту рамки."""
    widths = [int(g.get('w')) / E for g in tbl.find('a:tblGrid', NS).findall('a:gridCol', NS)]
    trs = tbl.findall('a:tr', NS)
    texts = [[''.join(t.text or '' for t in tc.iter('{%s}t' % NS['a'])) for tc in tr.findall('a:tc', NS)] for tr in trs]
    pad_v, pad_h = 11.4, 22.6  # поля ячейки шаблона: 0,2 см сверху и снизу, 0,4 см по бокам
    size = capped('таблица', start)
    while True:
        heights = [max([fitter.lines(t, size, max(10, w - pad_h)) * size * 1.2 + pad_v
                        for t, w in zip(row, widths)] + [0]) for row in texts]
        if sum(heights) <= avail or size <= floor:
            break
        size -= 0.5
    note_size('таблица', start, size)
    for tr, h in zip(trs, heights):
        tr.set('h', str(int(h * E)))
        for rpr in tr.iter('{%s}rPr' % NS['a'], '{%s}endParaRPr' % NS['a']):
            rpr.set('sz', str(int(size * 100)))
    if size < start:
        report.append(f'{where}: кегль таблицы {start:g} → {size:g} pt')
    if sum(heights) > avail:
        report.append(f'{where}: ПЕРЕПОЛНЕНИЕ — таблица выше рамки на {sum(heights) - avail:.0f} pt. Сократите текст или разбейте слайд')


def cell_text(tc, value):
    txBody = tc.find('a:txBody', NS)
    ps = txBody.findall('a:p', NS)
    p = ps[0]
    for extra in ps[1:]:
        txBody.remove(extra)
    runs = p.findall('a:r', NS)
    rPr = copy.deepcopy(runs[0].find('a:rPr', NS)) if runs else None
    for r in runs:
        p.remove(r)
    end = p.find('a:endParaRPr', NS)
    r = etree.Element('{%s}r' % NS['a'])
    if rPr is not None:
        r.append(rPr)
    t = etree.SubElement(r, '{%s}t' % NS['a'])
    t.text = nbsp(str(value))
    if end is not None:
        end.addprevious(r)
    else:
        p.append(r)


def fill_chart(gf, data):
    cd = CategoryChartData()
    cd.categories = data['categories']
    for s in data['series']:
        cd.add_series(s.get('name', ''), s['values'])
    gf.chart.replace_data(cd)
    # выноски-линии, нарисованные поверх демонстрационных данных, к новым данным не подходят
    cs = gf.chart.part._element
    for us in cs.findall('{http://schemas.openxmlformats.org/drawingml/2006/chart}userShapes'):
        gf.chart.part.drop_rel(us.get('{%s}id' % R_NS))
        cs.remove(us)


# ---------------------------------------------------------------- слайд
def fill_slide(prs, slide, spec, s, ctx, where):
    report, fitter, lib = ctx['report'], ctx['fitter'], ctx['lib']
    used = set()

    def put(idx, value, name, max_lines=None):
        sh = ph_by_idx(slide, idx)
        if sh is None:
            return
        used.add(idx)
        set_text(sh, value, slide, fitter, report, f'{where}/{name}', max_lines)

    # поля
    for name, idx in spec.get('fields', {}).items():
        if s.get(name) not in (None, ''):
            lines = None
            if name == 'title' and spec['proto'] > 14 and spec['proto'] != 52:
                lines = 2
            elif name.endswith('value'):
                lines = 1  # крупная цифра не переносится, а уменьшается
            put(idx, s[name], name, lines)
    # текстовые блоки, которые не плейсхолдеры
    for name, prefix in spec.get('textbox', {}).items():
        sh = find_shape(slide, prefix)
        if sh is None:
            continue
        if s.get(name):
            bp = sh.text_frame._txBody.find('a:bodyPr', NS)
            bp.set('wrap', 'none')  # короткое значение («45 %») не переносим
            set_text(sh, s[name], slide, fitter, report, f'{where}/{name}', 1)
            # длинное значение сдвигает соседнюю подпись вправо
            x, y, w, h = bbox(sh)
            need = fitter.width(sh.text_frame.text, eff_size(sh, slide)) * 1.1 + 12
            if need > w:
                for o in slide.shapes:
                    ox, oy, ow, oh = bbox(o) if o.width is not None else (0, 0, 0, 0)
                    if o is not sh and x + w - 2 <= ox <= x + w + 40 and oy < y + h and y < oy + oh:
                        # у плейсхолдера без своей рамки задаём все четыре координаты, иначе top станет 0
                        o.left, o.top = int((ox + need - w) * E), int(oy * E)
                        o.width, o.height = int(max(20, ow - (need - w)) * E), int(oh * E)
        else:
            delete(sh)
    for prefix in spec.get('drop', []):
        for sh in [x for x in slide.shapes if x.name.startswith(prefix)]:
            delete(sh)

    # элементы
    items = spec.get('items')
    if items:
        given = s.get('items') or []
        if len(given) > len(items):
            report.append(f'{where}: элементов {len(given)}, в макете {len(items)} — лишние отброшены')
        value_proto = next((i['value'] for i in items if i.get('value')), None)
        for k, it in enumerate(items):
            box = it['box']
            if k >= len(given):
                for sh in [x for x in slide.shapes if center_in(x, box)]:
                    delete(sh)
                continue
            g = given[k]
            if isinstance(g, str):
                g = {'text': g}
            for field in ('title', 'text', 'label', 'num', 'page', 'year', 'name', 'role', 'phone', 'email', 'qr_text'):
                if it.get(field) is None:
                    continue
                v = g.get(field)
                if field == 'num' and v is None:
                    v = f'{k + 1:02d}'
                if v not in (None, ''):
                    one = field in ('num', 'year', 'label', 'page', 'phone', 'email')  # не переносить, а уменьшать
                    put(it[field], v, f'{field}{k + 1}', 1 if one else None)
            # цифра: своя ячейка или копия ячейки соседа
            if 'value' in it:
                if g.get('value') not in (None, ''):
                    idx = it['value'] or value_proto
                    if it['value']:
                        put(idx, g['value'], f'value{k + 1}', 1)
                    else:
                        borrow_value(slide, items, idx, box, g['value'], ctx, f'{where}/value{k + 1}')
            # иконка: заменить, убрать или оставить из прототипа
            icons = [x for x in slide.shapes if not x.is_placeholder and stype(x) in (6, 13)
                     and center_in(x, box) and max(bbox(x)[2:]) <= 62 and not x.name.startswith('Иконка')]
            if it.get('icon'):
                want = g.get('icon')
                has_value = g.get('value') not in (None, '')
                for ic in icons:
                    if want:
                        t = bbox(ic)
                        side = max(t[2], t[3])
                        target = (t[0] + t[2] / 2 - side / 2, t[1] + t[3] / 2 - side / 2, side, side)
                        if place_icon(slide, lib, want, target, prs):
                            delete(ic)
                        else:
                            report.append(f'{where}: иконки «{want}» нет в библиотеке (см. references/icons.png)')
                    elif has_value or want is False:
                        delete(ic)
            # фото внутри элемента
            if it.get('photo'):
                pics = [x for x in slide.shapes if stype(x) == 13 and center_in(x, box) and bbox(x)[2] > 100]
                path = asset('photos', g.get('photo'))
                for pic in pics:
                    if path:
                        put_image(slide, pic, path)
                    elif g.get('photo'):
                        report.append(f'{where}: фото «{g["photo"]}» не найдено')
            # контакты: QR
            if it.get('qr'):
                qr = find_shape(slide, it['qr'])
                path = g.get('qr')
                if qr is not None:
                    if path and os.path.isfile(path):
                        x, y, w, h = bbox(qr)
                        slide.shapes.add_picture(path, Emu(int(x * E)), Emu(int(y * E)), Emu(int(w * E)), Emu(int(h * E)))
                    delete(qr)
                if not path:
                    widen_without_qr(slide, it['qr_box'])
                    sh = ph_by_idx(slide, it['qr_text'])
                    if sh is not None:
                        delete(sh)
                    if it.get('qr_panel'):
                        pn = find_shape(slide, it['qr_panel'])
                        if pn is not None:
                            delete(pn)

    # фото слайда
    if spec.get('photo'):
        pic = find_shape(slide, spec['photo'])
        path = asset('photos', s.get('photo'))
        if pic is not None and path:
            put_image(slide, pic, path)
        elif s.get('photo'):
            report.append(f'{where}: фото «{s["photo"]}» не найдено')
    # портрет на обложке
    if spec.get('portrait'):
        pic = find_shape(slide, spec['portrait'])
        if pic is not None and s.get('portrait', True) is not False:
            keep_text_off_portrait(slide, pic, fitter, report, where)
        p = s.get('portrait', ctx['direction'])
        if pic is not None:
            if p is False:
                delete(pic)
            else:
                path = asset('people', person_key(p))
                if path:
                    put_image(slide, pic, path, focus=(0.5, 0.0))
                else:
                    report.append(f'{where}: портрет «{p}» не найден')
    # растры направления: фон раздела, вертикальная карточка, карточка контакта
    n = ctx['dir_no']
    if spec.get('background'):
        pic = find_shape(slide, spec['background'])
        bg = asset('backgrounds', s.get('background') or f'01_{n:02d}')
        if pic is not None and bg:
            swap_blob(slide, pic, bg)
    if spec.get('highlight'):
        pic = find_shape(slide, spec['highlight'])
        if pic is not None:
            swap_blob(slide, pic, asset('cards', f'highlight_{n:02d}'))
    if spec.get('card'):
        pic = find_shape(slide, spec['card'])
        if pic is not None:
            swap_blob(slide, pic, asset('cards', f'card_{n:02d}'))
    # светлые направления (B2B, инфраструктура): белый текст на светлом фоне не читается
    if n in LIGHT_DIRECTIONS:
        for key in ('background', 'highlight', 'card'):
            pic = find_shape(slide, spec[key]) if spec.get(key) else None
            if pic is None:
                continue
            for sh in slide.shapes:
                if sh is not pic and sh.has_text_frame and center_in(sh, bbox(pic)):
                    for para in sh.text_frame.paragraphs:
                        for r in para.runs:
                            r.font.color.rgb = RGBColor.from_string(DARK_TEXT)
    # таблица
    if spec.get('table'):
        gf = next((x for x in slide.shapes if x.has_table), None)
        ob = spec.get('optional_box')
        if ob and not s.get('conclusion') and not s.get('conclusion_title'):
            for sh in [x for x in slide.shapes if center_in(x, ob)]:
                delete(sh)
            if gf is not None:  # без плашки-вывода таблица занимает всю высоту до низа контента
                gf.height = int((500 - gf.top / E) * E)
        if gf is not None and s.get('table'):
            fill_table(gf, s['table'], report, where, fitter)
    # диаграммы
    for k, name in enumerate(spec.get('charts', [])):
        gf = find_shape(slide, name)
        data = (s.get('charts') or [])[k] if k < len(s.get('charts') or []) else None
        if gf is not None and data:
            fill_chart(gf, data)
        elif gf is not None:
            report.append(f'{where}: диаграмма {k + 1} без данных — осталась демонстрационная, замените')
    # пустые плейсхолдеры убираем, чтобы в PowerPoint не висело «Образец текста»
    for sh in list(slide.placeholders):
        if sh.placeholder_format.type == 13:  # номер слайда
            continue
        if sh.placeholder_format.idx not in used and sh.has_text_frame:
            delete(sh)
    if s.get('notes'):
        set_notes(slide, s['notes'])


def keep_text_off_portrait(slide, pic, fitter, report, where):
    """Подзаголовок обложки не должен уходить под портрет: сужаем поле и даём ему вторую строку."""
    edge = pic.left / E + 40  # слева у вырезанного портрета прозрачное поле
    for sh in slide.placeholders:
        if sh.placeholder_format.type in (1, 3) or not sh.has_text_frame or not sh.text_frame.text.strip():
            continue
        x, y, w, h = bbox(sh)
        if x < edge < x + w and y + h > pic.top / E and y < 300:
            continue  # поле в верхнем углу (meta) — не трогаем
        if x < edge < x + w:
            size = eff_size(sh, slide)
            sh.left, sh.top, sh.width = int(x * E), int(y * E), int((edge - x - 8) * E)
            sh.height = int(max(h, size * 2.6) * E)
            fit(sh, slide, fitter, report, where + '/подзаголовок', 2)


def fill_house(slide, spec, s, ctx, where):
    """Макеты фирменного стиля команды (scripts/house.py)."""
    if spec['house'] in ('title', 'chapter'):
        pic = find_shape(slide, 'Image')
        p = s.get('portrait', ctx['direction']) if spec['house'] == 'title' else False
        if pic is not None:
            if p is False:
                delete(pic)
            else:
                path = asset('people', person_key(p))
                if path:
                    put_image(slide, pic, path, focus=(0.5, 0.0))
                else:
                    ctx['report'].append(f'{where}: портрет «{p}» не найден')
    house.HOUSE[spec['house']](slide, s, ctx, where)
    if s.get('notes'):
        set_notes(slide, s['notes'])


def set_notes(slide, text):
    """Заметки докладчика. В мастере заметок шаблона нет поля текста — добавляем его сами."""
    ns = slide.notes_slide
    if ns.notes_text_frame is None:
        sp = parse_xml(
            '<p:sp xmlns:p="%s" xmlns:a="%s"><p:nvSpPr><p:cNvPr id="%d" name="Заметки"/>'
            '<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr><p:ph type="body" idx="1"/></p:nvPr></p:nvSpPr>'
            '<p:spPr><a:xfrm><a:off x="685800" y="4400550"/><a:ext cx="5486400" cy="3600450"/></a:xfrm>'
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>'
            '<p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody></p:sp>'
            % (NS['p'], NS['a'], next_id(ns)))
        ns.shapes._spTree.append(sp)
    ns.notes_text_frame.text = text


def widen_without_qr(slide, qr_box):
    """Без QR-кода: убираем его подложку и растягиваем карточку контакта на всю ширину."""
    right = qr_box[0] + qr_box[2]
    for sh in list(slide.shapes):
        if stype(sh) == 6:
            for ch in list(sh.shapes):
                if center_in(ch, qr_box):
                    delete(ch)
            continue
        x, y, w, h = bbox(sh)
        if not sh.is_placeholder and y + h / 2 > qr_box[1] and y + h / 2 < qr_box[1] + qr_box[3]:
            if center_in(sh, qr_box):
                continue
            if qr_box[0] - 12 < x + w < qr_box[0] + 2:
                sh.width = int((right - x) * E)
    for sh in list(slide.shapes):  # подложки внутри групп
        if stype(sh) == 6:
            for ch in sh.shapes:
                x, y, w, h = bbox(ch)
                if w > 300 and center_in(ch, (0, qr_box[1], 960, qr_box[3])):
                    ch.width = int((right - x) * E)


def borrow_value(slide, items, idx, box, value, ctx, where):
    src = ph_by_idx(slide, idx)
    lay = next((p for p in slide.slide_layout.placeholders if p.placeholder_format.idx == idx), None)
    if src is None:
        return
    src_box = next(i['box'] for i in items if i.get('value') == idx)
    el = copy.deepcopy(src._element)
    off = xp(el, './p:spPr/a:xfrm/a:off')
    if not off and lay is not None:
        el.find('p:spPr', NS).append(copy.deepcopy(xp(lay._element, './p:spPr/a:xfrm')[0]))
        off = xp(el, './p:spPr/a:xfrm/a:off')
    off = off[0]
    off.set('x', str(int(off.get('x')) + int((box[0] - src_box[0]) * E)))
    off.set('y', str(int(off.get('y')) + int((box[1] - src_box[1]) * E)))
    # отвязываем от плейсхолдера макета: переносим стиль макета внутрь фигуры
    ph = xp(el, './p:nvSpPr/p:nvPr/p:ph')[0]
    ph.getparent().remove(ph)
    if lay is not None:
        ls = xp(lay._element, './p:txBody/a:lstStyle')
        mine = xp(el, './p:txBody/a:lstStyle')
        if ls and mine:
            mine[0].getparent().replace(mine[0], copy.deepcopy(ls[0]))
        bp = xp(lay._element, './p:txBody/a:bodyPr')
        mb = xp(el, './p:txBody/a:bodyPr')
        if bp and mb:
            mb[0].getparent().replace(mb[0], copy.deepcopy(bp[0]))
    slide.shapes._spTree.append(el)
    renumber(slide)
    sh = slide.shapes[-1]
    set_text(sh, value, slide, ctx['fitter'], ctx['report'], where, 1)


# ---------------------------------------------------------------- партнёр и логотип на обложках
def fix_cover_layout(layout, partner):
    shapes = list(layout.shapes)
    partner_shapes = [sh for sh in shapes if not sh.is_placeholder and sh.top / E < 60 and 100 < sh.left / E < 260
                      and sh.width / E < 100]
    short_logo = [sh for sh in shapes if not sh.is_placeholder and stype(sh) == 13 and sh.top / E < 60
                  and sh.left / E < 40 and 60 < sh.width / E < 80]
    if not partner_shapes:
        return
    if partner:
        pic = [sh for sh in partner_shapes if sh.width / E > 20]
        for sh in pic:
            img_part, rId = layout.part.get_or_add_image_part(partner)
            blip = xp(sh._element, './/a:blip')[0]
            blip.set('{%s}embed' % R_NS, rId)
            for ext in blip.findall('a:extLst', NS):
                blip.remove(ext)
            with Image.open(partner) as im:
                w, h = im.size
            hh = sh.height
            sh.width = int(hh * w / h)
            sh.crop_left = sh.crop_right = sh.crop_top = sh.crop_bottom = 0
        return
    for sh in partner_shapes:
        delete(sh)
    for sh in short_logo:  # без партнёра — полный логотип «ВТБ Образование»
        img_part, rId = layout.part.get_or_add_image_part(os.path.join(ASSETS, 'logo', 'logo-color.png'))
        blip = xp(sh._element, './/a:blip')[0]
        blip.set('{%s}embed' % R_NS, rId)
        for ext in blip.findall('a:extLst', NS):
            blip.remove(ext)
        sh.width = int(sh.height * 1374 / 240)
        sh.crop_left = sh.crop_right = sh.crop_top = sh.crop_bottom = 0


# ---------------------------------------------------------------- сборка
def norm(name):
    return re.sub(r'^\d+_', '', name.strip())


def drop_slide(prs, slide):
    sld = prs.slides._sldIdLst
    for sid in list(sld):
        if prs.part.related_part(sid.rId) is slide.part:
            prs.part.drop_rel(sid.rId)
            sld.remove(sid)
            return


def restyle(prs):
    """Шрифт VTB Group → VTB Omega UI и отказ от прописных (cap="all") во всём пакете: темы, мастера,
    макеты, слайды, диаграммы. Текст набирается так, как написан."""
    pat = re.compile(rb'typeface="(VTB Group(?: Book| Light| Bold| Medium| DemiBold)?)"')
    for part in prs.part.package.iter_parts():
        if hasattr(part, '_element'):
            for el in part._element.iter():
                tf = el.get('typeface')
                if tf in FONT_MAP:
                    el.set('typeface', FONT_MAP[tf])
                    for a in ('panose', 'pitchFamily', 'charset'):
                        el.attrib.pop(a, None)
                if el.get('cap') in ('all', 'small'):
                    del el.attrib['cap']
        elif part.content_type.endswith('theme+xml'):
            blob = pat.sub(b'typeface="' + FONT.encode() + b'"', part.blob)
            # шрифты темы (заголовки и текст) во всех темах, включая заметки и раздаточные материалы
            blob = re.sub(rb'(<a:(?:major|minor)Font>\s*<a:latin typeface=")[^"]*(")',
                          rb'\1' + FONT.encode() + rb'\2', blob)
            part._blob = blob


def build(deck, out_dir, template):
    """Собирает колоду; если какие-то элементы пришлось уменьшить, пересобирает с единым кеглем этих
    элементов на всех слайдах (до 3 проходов)."""
    UNIFORM['caps'] = {}
    for _ in range(3):
        UNIFORM['shrunk'] = {}
        prs, report = build_once(deck, template)
        merged = dict(UNIFORM['caps'])
        for role, v in UNIFORM['shrunk'].items():
            merged[role] = min(v, merged.get(role, v))
        if merged == UNIFORM['caps']:
            break
        UNIFORM['caps'] = merged
    if UNIFORM['caps']:
        report.append('единый кегль по колоде: ' + ', '.join(f'{r} — {v:g} pt' for r, v in sorted(UNIFORM['caps'].items())))
    os.makedirs(out_dir, exist_ok=True)
    fn = os.path.join(out_dir, deck.get('file', 'deck') + '.pptx')
    prs.save(fn)
    return fn, report


def build_once(deck, template):
    prs = Presentation(template)
    restyle(prs)
    direction = deck.get('direction', 'vtb')
    if direction not in DIRECTIONS:
        raise SystemExit(f'Неизвестное направление «{direction}». Есть: {", ".join(DIRECTIONS)}')
    n = DIRECTIONS[direction][0]
    master = prs.slide_masters[n - 1]
    protos = {s._element.cSld.get('name'): s for s in prs.slides}
    originals = list(prs.slides)
    ctx = dict(report=[], fitter=Fitter(), lib=build_library(prs), direction=direction, dir_no=n, prs=prs)
    exact = {l.name.strip(): l for l in master.slide_layouts}
    loose = {norm(l.name): l for l in master.slide_layouts}

    partner = deck.get('partner')
    if partner and not os.path.isfile(partner):
        ctx['report'].append(f'логотип партнёра «{partner}» не найден — обложка без партнёра')
        partner = None
    fixed = set()
    for i, s in enumerate(deck['slides'], 1):
        name = s.get('layout')
        if name not in LAYOUTS:
            raise SystemExit(f'Слайд {i}: нет макета «{name}». Есть: {", ".join(LAYOUTS)}')
        spec = LAYOUTS[name]
        proto = protos[f'proto-{spec["proto"]}']
        pname = proto.slide_layout.name.strip()
        layout = exact.get(pname) or loose[norm(pname)]
        if spec.get('partner') or spec['proto'] in (9, 10, 11, 12, 13, 14):
            if id(layout) not in fixed:
                fix_cover_layout(layout, partner)
                fixed.add(id(layout))
        slide = clone_slide(prs, proto, layout)
        where = f'слайд {i} ({name})'
        if spec.get('house'):
            fill_house(slide, spec, s, ctx, where)
        else:
            fill_slide(prs, slide, spec, s, ctx, where)

    for sl in originals:
        drop_slide(prs, sl)
    # оставляем один мастер и только использованные макеты
    used_layouts = {id(s.slide_layout) for s in prs.slides}
    ids = prs.part._element.find('p:sldMasterIdLst', NS)
    for m_i, (mid, m) in enumerate(zip(list(ids), list(prs.slide_masters))):
        if m is not master:
            prs.part.drop_rel(mid.get('{%s}id' % R_NS))
            ids.remove(mid)
    for lay in list(master.slide_layouts):
        if id(lay) not in used_layouts:
            master.slide_layouts.remove(lay)

    return prs, ctx['report']


def preview(pptx, out_dir, prefix='slide', labels=None):
    if not shutil.which('soffice'):
        print('LibreOffice не найден — превью пропущено')
        return
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', out_dir, pptx],
                   capture_output=True, timeout=600)
    pdf = os.path.splitext(pptx)[0] + '.pdf'
    if shutil.which('pdftoppm') and os.path.exists(pdf):
        for old in os.listdir(out_dir):
            if re.fullmatch(prefix + r'-\d+\.png', old):
                os.remove(os.path.join(out_dir, old))
        subprocess.run(['pdftoppm', '-r', '80', '-png', pdf, os.path.join(out_dir, prefix)])
        sheet = contact_sheet(out_dir, prefix, labels)
        print('превью:', pdf, f'и {prefix}-*.png; все слайды на одном листе — {sheet}')


def contact_sheet(out_dir, prefix, labels=None, cols=4):
    """Все слайды на одном листе с подписями «N · макет» — для согласования и визуального ревью."""
    from PIL import ImageDraw
    files = sorted(f for f in os.listdir(out_dir) if re.fullmatch(prefix + r'-\d+\.png', f))
    if not files:
        return None
    ims = [Image.open(os.path.join(out_dir, f)).convert('RGB') for f in files]
    tw = 480
    th = int(ims[0].height * tw / ims[0].width)
    g, lab = 12, 24
    rows = (len(ims) + cols - 1) // cols
    sheet = Image.new('RGB', (cols * tw + (cols + 1) * g, rows * (th + lab) + (rows + 1) * g), 'white')
    d = ImageDraw.Draw(sheet)
    try:
        font = ImageFont.truetype(Fitter().font_path, 16)
    except OSError:
        font = None
    for i, im in enumerate(ims):
        x, y = g + (i % cols) * (tw + g), g + (i // cols) * (th + lab + g)
        name = labels[i] if labels and i < len(labels) else ''
        d.text((x, y + 2), f'{i + 1} · {name}', fill=(10, 40, 150), font=font)
        sheet.paste(im.resize((tw, th), Image.LANCZOS), (x, y + lab))
        d.rectangle([x, y + lab, x + tw - 1, y + lab + th - 1], outline=(210, 214, 222))
    path = os.path.join(out_dir, prefix + '-sheet.png')
    sheet.save(path)
    return path


def list_people():
    print('Портреты для обложек ("portrait"):')
    for key, (ru, who, direction) in PEOPLE.items():
        d = f'по умолчанию для {direction}' if direction else 'без направления'
        print(f'  {key:<8} {ru:<8} {who} — {d}')


def list_icons(template):
    lib = build_library(Presentation(template))
    icons = sorted(k for k in lib if not k.startswith('ill-'))
    ills = sorted(k for k in lib if k.startswith('ill-'))
    print('Иконки:', ', '.join(icons))
    print('Иллюстрации:', ', '.join(ills))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('deck', nargs='?')
    ap.add_argument('-o', '--out', default='out')
    ap.add_argument('--template', default=os.path.join(ASSETS, 'template.pptx'))
    ap.add_argument('--preview', action='store_true')
    ap.add_argument('--wireframe', action='store_true',
                    help='прототип: раскладка блоков в сером для согласования (дизайн — после «да»)')
    ap.add_argument('--no-review', action='store_true', help='не запускать ревьюер после сборки')
    ap.add_argument('--list-icons', action='store_true')
    ap.add_argument('--list-people', action='store_true')
    a = ap.parse_args()
    if a.list_people:
        return list_people()
    if a.list_icons:
        return list_icons(a.template)
    if not a.deck:
        ap.error('нужен deck.json')
    with open(a.deck, encoding='utf-8') as f:
        deck = json.load(f)
    base = os.path.dirname(os.path.abspath(a.deck))
    cwd = os.getcwd()
    os.chdir(base)  # относительные пути картинок — от deck.json
    out_dir = os.path.join(cwd, a.out) if not os.path.isabs(a.out) else a.out
    template = os.path.join(cwd, a.template) if not os.path.isabs(a.template) else a.template
    fn, report = build(deck, out_dir, template)
    os.chdir(cwd)
    labels = [s.get('layout', '') for s in deck['slides']]
    if a.wireframe:
        import wireframe
        proto = os.path.splitext(fn)[0] + '-прототип.pptx'
        wireframe.convert(Presentation(fn), deck).save(proto)
        os.remove(fn)  # дизайн собирается только после согласования прототипа
        print('прототип:', proto)
        for r in report:
            print(' -', r)
        preview(proto, out_dir, 'proto', labels)
        return 2 if any('ПЕРЕПОЛНЕНИЕ' in r for r in report) else 0
    print('готово:', fn)
    with open(os.path.join(out_dir, 'report.json'), 'w', encoding='utf-8') as f:
        json.dump(report, f, ensure_ascii=False, indent=1)
    for r in report:
        print(' -', r)
    errors = 0
    if not a.no_review:  # ревьюер: обязательная проверка перед отдачей
        import review
        prs = Presentation(fn)
        fixed = review.fix_hanging(prs)
        if sum(fixed.values()):
            prs.save(fn)
        rv = review.Review(Presentation(fn)).run()
        print('ревью (scripts/review.py):')
        print(rv.report(+fixed, True))
        errors = sum(1 for i in rv.items if i['level'] == 'ОШИБКА')
    if a.preview:
        preview(fn, out_dir, 'slide', labels)
    return 2 if any('ПЕРЕПОЛНЕНИЕ' in r for r in report) or errors else 0


if __name__ == '__main__':
    sys.exit(main())
