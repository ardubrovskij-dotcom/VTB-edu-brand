"""Прототип колоды: та же раскладка, что у дизайна, но в сером — чтобы согласовать расстановку блоков до дизайна.

Собирается из готовой колоды (build_deck.py --wireframe): фон белый, логотип и фоны мастера скрыты, карточки —
светло-серые, акцентные блоки и градиенты — серые, фото и портреты — серые заглушки с подписью «фото»,
иконки — серые квадраты, текст настоящий (тёмно-серый). В углу — номер слайда и имя макета.
"""
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.util import Emu, Pt

from build_deck import stype

E = 12700
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
INK = '2B2B2B'
CARD = 'EEF0F4'
ACCENT = 'C9CED8'
PHOTO = 'DADDE3'
LINE = 'B4BAC6'


def _solid(sp_pr, colour):
    for tag in ('solidFill', 'gradFill', 'blipFill', 'pattFill', 'noFill'):
        for el in sp_pr.findall('{%s}%s' % (A, tag)):
            sp_pr.remove(el)
    fill = parse_xml(f'<a:solidFill xmlns:a="{A}"><a:srgbClr val="{colour}"/></a:solidFill>')
    geom = sp_pr.find('{%s}prstGeom' % A)
    if geom is None:
        geom = sp_pr.find('{%s}custGeom' % A)
    if geom is not None:
        geom.addnext(fill)
    else:
        sp_pr.append(fill)


def _stub(slide, el, x, y, w, h, label, colour=PHOTO, outline=True):
    """Серая заглушка на месте картинки: тот же слой (текст поверх остаётся поверх)."""
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(colour)
    if outline:
        sh.line.color.rgb = RGBColor.from_string(LINE)
        sh.line.width = Pt(0.75)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    sh.name = f'Прототип: {label}'
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    if label and min(w, h) / E >= 40:
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = label
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor.from_string('7A8291')
    el.addprevious(sh._element)
    el.getparent().remove(el)


def _silhouette(slide, pic):
    """Портрет-вырезка → серый силуэт по прозрачности: видно, где фигура, но без фото."""
    import io
    from PIL import Image
    try:
        im = Image.open(io.BytesIO(pic.image.blob))
    except Exception:
        return False
    if im.mode not in ('RGBA', 'LA', 'P'):
        return False
    a = im.convert('RGBA').getchannel('A')
    if a.getextrema()[0] > 200:
        return False
    sil = Image.new('RGBA', im.size, (200, 205, 214, 0))
    sil.putalpha(a)
    buf = io.BytesIO()
    sil.save(buf, 'PNG')
    buf.seek(0)
    _, rid = slide.part.get_or_add_image_part(buf)
    pic._element.blipFill.blip.rEmbed = rid
    return True


def _texts(el):
    for rpr in el.iter('{%s}rPr' % A, '{%s}endParaRPr' % A, '{%s}defRPr' % A):
        for f in rpr.findall('{%s}solidFill' % A) + rpr.findall('{%s}gradFill' % A):
            rpr.remove(f)
        rpr.insert(0, parse_xml(f'<a:solidFill xmlns:a="{A}"><a:srgbClr val="{INK}"/></a:solidFill>'))
    for bu in el.iter('{%s}buClr' % A):
        for c in list(bu):
            bu.remove(c)
        bu.append(parse_xml(f'<a:srgbClr xmlns:a="{A}" val="{INK}"/>'))


def convert(prs, deck):
    names = [s.get('layout', '') for s in deck.get('slides', [])]
    for no, slide in enumerate(prs.slides, 1):
        sld = slide._element
        sld.set('showMasterSp', '0')  # без логотипа, фонов и линий мастера
        csld = sld.find('{%s}cSld' % P)
        bg = csld.find('{%s}bg' % P)
        if bg is not None:
            csld.remove(bg)
        csld.insert(0, parse_xml(f'<p:bg xmlns:p="{P}" xmlns:a="{A}"><p:bgPr><a:solidFill><a:srgbClr val="FFFFFF"/>'
                                 '</a:solidFill><a:effectLst/></p:bgPr></p:bg>'))
        for sh in list(slide.shapes):
            el = sh._element
            x, y, w, h = sh.left or 0, sh.top or 0, sh.width or 0, sh.height or 0
            if sh.is_placeholder and sh.placeholder_format.type == 13:
                el.getparent().remove(el)  # номер слайда — в подписи прототипа
                continue
            if el.xpath('.//a:alphaModFix') or sh.name.startswith('Водяной знак'):  # в прототипе не нужен
                el.getparent().remove(el)
                continue
            label = 'фото'
            if stype(sh) == 13 and max(w, h) / E > 100 and _silhouette(slide, sh):
                continue  # вырезка (портрет) — серый силуэт той же формы
            if stype(sh) in (13, 6):  # заглушка не выходит за слайд
                x1, y1 = min(x + w, prs.slide_width), min(y + h, prs.slide_height)
                x, y = max(0, x), max(0, y)
                w, h = x1 - x, y1 - y
            if stype(sh) == 13:  # картинка
                small = max(w, h) / E <= 100
                _stub(slide, el, x, y, w, h, '' if small else label, ACCENT if small else PHOTO)
                continue
            if stype(sh) == 6 and max(w, h) / E <= 100:  # иконка или водяной знак
                _stub(slide, el, x, y, w, h, '', ACCENT, outline=False)
                continue
            sp_pr = el.find('{%s}spPr' % P)
            if sp_pr is not None and stype(sh) != 6:
                if sp_pr.find('{%s}gradFill' % A) is not None or sp_pr.find('{%s}blipFill' % A) is not None:
                    _solid(sp_pr, ACCENT)
                elif sp_pr.find('{%s}solidFill' % A) is not None:
                    _solid(sp_pr, CARD)
            if getattr(sh, 'has_table', False) and sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells:
                        tc_pr = c._tc.get_or_add_tcPr()
                        for f in tc_pr.findall('{%s}gradFill' % A) + tc_pr.findall('{%s}solidFill' % A):
                            tc_pr.remove(f)
                        tc_pr.append(parse_xml(f'<a:solidFill xmlns:a="{A}"><a:srgbClr val="{CARD}"/></a:solidFill>'))
            _texts(el)
        lab = slide.shapes.add_textbox(Emu(30 * E), Emu(518 * E), Emu(600 * E), Emu(14 * E))
        r = lab.text_frame.paragraphs[0].add_run()
        r.text = f'прототип · слайд {no} · макет {names[no - 1] if no <= len(names) else ""}'
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor.from_string('9AA1AD')
        lab.name = 'Прототип: подпись'
    return prs
