#!/usr/bin/env python3
"""Переводит готовую презентацию на шрифт VTB Omega UI и убирает прописные.

Для файлов, собранных раньше или доработанных вручную в PowerPoint: правки в файле сохраняются.
    python3 tools/restyle_pptx.py "старая.pptx" -o "новая.pptx"

Что меняется:
- шрифты темы (заголовки и текст) во всех темах — VTB Omega UI;
- VTB Group Book / Light → VTB Omega UI, VTB Group Bold / Medium → VTB Omega UI Medium;
- атрибут «все прописные» (cap="all") снимается во всех мастерах, макетах и слайдах;
- текст, набранный ПРОПИСНЫМИ вручную (от двух слов), переводится в обычный регистр; аббревиатуры
  (ВТБ, РФ, ШОиФД, МСФО…) сохраняются. Список — ACRONYMS.
"""
import argparse
import os
import re
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'scripts'))
from pptx import Presentation  # noqa: E402
import build_deck  # noqa: E402

ACRONYMS = ['ВТБ', 'РФ', 'ШОиФД', 'МСФО', 'ДПО', 'АНО', 'ДО', 'ЦФА', 'LMS', 'KPI', 'ДВОПиИ', 'IFRS', 'IAS', 'ECL',
            'ПАО', 'СПбГУ', 'ВШМ', 'АХР', 'КВ', 'НМА', 'НИОКР', 'B2B', 'TBD', 'QR', 'HR', 'ИТ', 'IT']
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'


# имена собственные из нескольких слов: после перевода в строчные возвращаем им заглавные
PROPER = ['ВТБ Образование', 'Финансовый университет', 'Финансового университета', 'Финансовым университетом',
          'Финансовый департамент', 'Финансового департамента', 'Профессиональные финансы', 'Банк ВТБ',
          'Москва', 'Москве', 'Санкт-Петербург', 'Санкт-Петербурге', 'Private Banking', 'Привилегия']
SMALL_WORDS = {'и', 'в', 'во', 'на', 'с', 'со', 'к', 'по', 'для', 'из', 'от', 'до', 'о', 'об', 'за', 'не', 'а'}


def sentence_case(text):
    words = re.findall(r'[A-Za-zА-Яа-яЁё-]+', text)
    if len(words) <= 3 and not any(w.lower() in SMALL_WORDS for w in words) and not re.search(r'\d', text):
        # 2–3 слова без предлогов — скорее всего имя («ЖУРАВЛЕВА ЕЛЕНА»): каждое слово с заглавной
        out = re.sub(r'[A-Za-zА-Яа-яЁё]+', lambda m: m.group(0).capitalize(), text.lower())
    else:
        low = text.lower()
        out = low[:1].upper() + low[1:]
    for acr in ACRONYMS:
        out = re.sub(r'(?<![\w])' + re.escape(acr.lower()) + r'(?![\w])', acr, out, flags=re.I)
    for name in PROPER:
        out = re.sub(re.escape(name).replace('\\ ', '[ \u00a0]'), name, out, flags=re.I)
    return out


def fix_typed_caps(prs):
    n = 0
    parts = list(prs.slides) + [l for m in prs.slide_masters for l in m.slide_layouts]
    for sl in parts:
        for p in sl.shapes._spTree.iter(A + 'p'):
            runs = list(p.iter(A + 't'))
            text = ''.join(t.text or '' for t in runs)
            letters = re.sub(r'[^A-Za-zА-Яа-яЁё]', '', text)
            words = re.findall(r'[A-Za-zА-Яа-яЁё]{2,}', text)
            if len(words) >= 2 and letters.isupper() and not all(w in ACRONYMS for w in words):
                new = sentence_case(text)
                pos = 0
                for t in runs:  # раскладываем обратно по исходным ранам, сохраняя их формат
                    k = len(t.text or '')
                    t.text = new[pos:pos + k]
                    pos += k
                n += 1
    return n


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('src')
    ap.add_argument('-o', '--out', required=True)
    a = ap.parse_args()
    prs = Presentation(a.src)
    build_deck.restyle(prs)
    n = fix_typed_caps(prs)
    prs.save(a.out)
    print(f'готово: {a.out}; абзацев из прописных переведено в обычный регистр: {n}')


if __name__ == '__main__':
    main()
