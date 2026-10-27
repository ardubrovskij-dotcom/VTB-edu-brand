#!/usr/bin/env python3
"""Готовит assets/template.pptx из исходного шаблона «мерико × ВТБ Образование».

Запускать, когда маркетинг обновил шаблон в папке «ВТБ О брендинг»:
    python3 tools/prepare_template.py "<исходный>.pptx" -o assets/template.pptx

Что делает:
- оставляет слайды-прототипы (9–49), раздел (52) и библиотеку иконок и иллюстраций (63–75);
- во всех 8 мастерах оставляет только макеты, на которых стоят эти слайды;
- пережимает картинки (JPEG до 1920 px, q85), большие PNG уменьшает;
- по умолчанию удаляет встроенный шрифт VTB Group (репозиторий публичный). --keep-fonts оставляет его.
"""
import argparse, io, re, sys
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

KEEP = set(range(9, 50)) | {52} | set(range(63, 76))
NS = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}


def norm(name):
    return re.sub(r'^\d+_', '', name.strip())


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('src')
    ap.add_argument('-o', '--out', required=True)
    ap.add_argument('--keep-fonts', action='store_true')
    a = ap.parse_args()

    prs = Presentation(a.src)
    sld = prs.slides._sldIdLst
    for i, (sid, slide) in enumerate(zip(list(sld), list(prs.slides)), 1):
        # номер слайда в исходном шаблоне — имя прототипа, по нему его находит build_deck.py
        slide._element.cSld.set('name', f'proto-{i}')
        if i not in KEEP:
            prs.part.drop_rel(sid.rId)
            sld.remove(sid)

    used = {norm(s.slide_layout.name) for s in prs.slides}
    for m in prs.slide_masters:
        for lay in list(m.slide_layouts):
            if norm(lay.name) not in used:
                m.slide_layouts.remove(lay)

    if not a.keep_fonts:
        root = prs.part._element
        lst = root.find('p:embeddedFontLst', NS)
        if lst is not None:
            for rid in lst.xpath('.//@r:id', namespaces=NS):
                prs.part.drop_rel(rid)
            root.remove(lst)
        root.set('embedTrueTypeFonts', '0')

    before = after = 0
    seen = set()
    for rel in [r for p in prs.part.package.iter_parts() for r in p.rels.values()]:
        if rel.is_external or rel.reltype != RT.IMAGE:
            continue
        part = rel.target_part
        if id(part) in seen:
            continue
        seen.add(id(part))
        blob = part.blob
        before += len(blob)
        ct = part.content_type
        try:
            im = Image.open(io.BytesIO(blob))
        except Exception:
            after += len(blob)
            continue
        out = None
        if ct == 'image/jpeg' and len(blob) > 300_000:
            im = im.convert('RGB')
            if im.width > 1920:
                im = im.resize((1920, round(im.height * 1920 / im.width)), Image.LANCZOS)
            b = io.BytesIO(); im.save(b, 'JPEG', quality=85, optimize=True); out = b.getvalue()
        elif ct == 'image/png' and len(blob) > 600_000:
            if im.width > 1400:
                im = im.resize((1400, round(im.height * 1400 / im.width)), Image.LANCZOS)
            b = io.BytesIO(); im.save(b, 'PNG', optimize=True); out = b.getvalue()
        if out and len(out) < len(blob):
            part._blob = out
            blob = out
        after += len(blob)
    prs.save(a.out)
    print(f'слайдов: {len(prs.slides)}; картинки {before/1e6:.1f} → {after/1e6:.1f} МБ; сохранено: {a.out}')


if __name__ == '__main__':
    sys.exit(main())
