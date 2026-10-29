#!/usr/bin/env python3
"""Самопроверка ревьюера: ломает собранную колоду-пример и проверяет, что review.py находит каждую поломку.

    python3 tools/test_review.py      # собирает examples/shoifd.json во временную папку
"""
import os, subprocess, sys, tempfile

from pptx import Presentation
from pptx.util import Pt, Emu

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(SKILL, 'scripts'))
import review  # noqa: E402

E = 12700


def by_name(slide, name):
    return [s for s in slide.shapes if s.name == name]


def main():
    tmp = tempfile.mkdtemp()
    subprocess.run([sys.executable, os.path.join(SKILL, 'scripts', 'build_deck.py'),
                    os.path.join(SKILL, 'examples', 'shoifd.json'), '-o', tmp], check=True, capture_output=True)
    src = os.path.join(tmp, 'shoifd.pptx')
    clean = review.Review(Presentation(src)).run()
    errors = [i for i in clean.items if i['level'] == 'ОШИБКА']
    assert not errors, 'эталонная сборка должна проходить без ошибок:\n' + '\n'.join(i['msg'] for i in errors)

    prs = Presentation(src)
    s = prs.slides
    expect = []
    # 1. текст налезает на текст: подпись цифры поднята на цифру
    cap = by_name(s[2], 'Подпись')[0]
    cap.top = by_name(s[2], 'Цифра')[0].top
    expect.append((3, 'наложение'))
    # 2. иконка искажена и сдвинута
    icons = [x for x in s[5].shapes if x.shape_type == 6 and x.width < 40 * E]
    icons[0].width = Emu(int(icons[0].width * 1.6))
    icons[1].left = Emu(icons[1].left + 9 * E)
    expect += [(6, 'иконка')]
    # 3. чужой шрифт и капс
    r = s[1].shapes.title.text_frame.paragraphs[0].runs[0]
    r.font.name = 'Arial'
    r.text = 'ИЗМЕНЕНИЯ ' + r.text
    expect += [(2, 'шрифт'), (2, 'капс')]
    # 4. автосжатие заголовка и разный кегль заголовков
    t = s[3].shapes.title.text_frame
    bp = t._txBody.bodyPr
    for el in list(bp):
        bp.remove(el)
    from pptx.oxml import parse_xml
    bp.append(parse_xml('<a:normAutofit xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" fontScale="90000"/>'))
    expect += [(4, 'кегль'), (0, 'кегль')]
    # 5. карточка ниже нижней границы
    card = by_name(s[1], 'Карточка')[0]
    card.top = Emu(card.top + 30 * E)
    expect.append((2, 'поля'))
    # 6. портрет обрезан рамкой: уже квадратного портрета рамка
    pic = [x for x in s[0].shapes if x.shape_type == 13][0]
    pic.crop_right = 0.25
    expect.append((1, 'фото'))
    # 7. растянутое фото: сцена из assets/photos в рамке не своих пропорций
    photo = os.path.join(SKILL, 'assets', 'photos', sorted(os.listdir(os.path.join(SKILL, 'assets', 'photos')))[0])
    s[8].shapes.add_picture(photo, Emu(600 * E), Emu(60 * E), Emu(300 * E), Emu(60 * E))
    expect.append((9, 'фото'))
    # 8. пустое поле и заголовок длиннее 3 строк
    s[8].shapes.add_textbox(Emu(200 * E), Emu(200 * E), Emu(50 * E), Emu(20 * E))
    s[8].shapes.title.text_frame.paragraphs[0].runs[0].text += ' и ещё очень длинный хвост заголовка, ' * 3
    expect += [(9, 'пусто'), (9, 'заголовок')]
    # 9. слишком мелкий текст
    for p in by_name(s[6], 'Текст')[-1].text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(7)
    expect.append((7, 'кегль'))

    # 10. водяной знак вылезает за строку на сноску; 11. обводка иконки не уменьшена вместе с иконкой
    mark = [x for x in s[7].shapes if x.shape_type == 6 and 40 * E < x.width <= 100 * E][-1]
    mark.top = Emu(mark.top + 40 * E)
    expect.append((8, 'иконка'))
    ic = [x for x in s[5].shapes if x.shape_type == 6 and x.width < 40 * E][2]
    for ln in ic._element.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}ln'):
        ln.set('w', str(int(3.6 * E)))
    if not list(ic._element.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}ln')):
        sp = ic._element.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}spPr')
        sp.append(parse_xml('<a:ln xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" w="45720"/>'))
    broken = os.path.join(tmp, 'broken.pptx')
    prs.save(broken)
    rv = review.Review(Presentation(broken)).run()
    found = {(i['slide'], i['kind']) for i in rv.items if i['level'] == 'ОШИБКА'}
    missed = [e for e in expect if e not in found]
    for e in expect:
        print(('OK    ' if e in found else 'ПРОПУЩЕНО ') + f'слайд {e[0]}: {e[1]}')
    # висячие предлоги: --fix ставит неразрывные пробелы
    prs = Presentation(broken)
    s0 = prs.slides[2].shapes.title.text_frame.paragraphs[0].runs[0]
    s0.text = 'Проект в банке и для студентов'
    n = review.fix_hanging(prs)
    assert '\xa0' in prs.slides[2].shapes.title.text_frame.text and n[3] >= 3, 'висячие предлоги не исправлены'
    print('OK    висячие предлоги исправлены:', n[3])
    if missed:
        sys.exit('ревьюер пропустил: ' + ', '.join(map(str, missed)))
    print('ревьюер находит все поломки')


if __name__ == '__main__':
    main()
